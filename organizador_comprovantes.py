#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

"""Organizador de Comprovantes (arquivo único).

ARQUITETURA (arquivo único):
1) Runtime/config/logging
2) Modelos e normalização
3) Extração PDF/OCR
4) Classificação e heurísticas
5) Relatórios (Excel/CSV)
6) Interface Tkinter (dashboard unificado)
7) Diagnóstico e build PyInstaller
8) Entrypoint principal
"""

import argparse
import json
import logging
import os
import re
import sqlite3
import shutil
import subprocess
import sys
import unicodedata
import threading
import time
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable

from app.config.settings import load_app_config
from app.core.text_model import TextoProcessado
from app.diagnostics.environment import run_environment_diagnostics
from app.infrastructure.paths import is_frozen as infra_is_frozen, app_root as infra_app_root, data_path as infra_data_path
from app.services.pipeline import ReceiptPipeline, PipelineDeps

APP_NOME = "Separador de Boletos / Organizador de Comprovantes"
APP_VERSAO = "2.1.0"
APP_CFG = load_app_config()
PDF_TEXT_CACHE: dict[str, tuple[float, str, str, bool]] = {}


BANCOS_CONHECIDOS = {
    "BANCO DO BRASIL", "CAIXA ECONOMICA FEDERAL", "BRADESCO", "ITAU", "ITAU UNIBANCO",
    "SANTANDER", "NUBANK", "BANCO INTER", "INTER", "C6 BANK", "BANCO ORIGINAL",
    "ORIGINAL", "NEXT", "BTG PACTUAL", "BTG", "SAFRA", "BANRISUL", "SICOOB", "SICREDI",
    "BANCO MASTER"
}

BANCOS_CODIGO = {
    "001": "BANCO DO BRASIL",
    "033": "SANTANDER",
    "104": "CAIXA ECONOMICA FEDERAL",
    "237": "BRADESCO",
    "341": "ITAU",
}

PREFIXOS_TECNICOS = ("BOL", "PIX", "TRANF", "DOC")

ROTULOS_CONTAMINANTES = [
    "NOME BENEFICIARIO", "RAZAO SOCIAL BENEFICIARIO", "BENEFICIARIO FINAL",
    "BANCO DESTINATARIO", "INSTITUICAO RECEBEDORA", "VALOR TOTAL",
    "VALOR DO PAGAMENTO", "VALOR PAGO", "DATA DO PAGAMENTO", "DATA PAGAMENTO",
    "DATA DO DEBITO", "LINHA DIGITAVEL", "CODIGO DE BARRAS", "AUTENTICACAO",
    "NOME DO PAGADOR", "HORA", "HORARIO"
]

ROTULOS_SEMANTICOS_POSITIVOS: dict[str, dict[str, object]] = {
    "direto_recebedor": {
        "prioridade": 130,
        "rotulos": [
            "nome do recebedor", "nome do beneficiario", "beneficiario", "beneficiario final",
            "recebedor", "destinatario", "favorecido", "favorecido final",
            "favorecido da transferencia", "destinatario do pagamento", "nome do favorecido",
            "nome do destinatario", "nome do beneficiario final", "nome do recebedor final",
            "beneficiario da operacao", "destinatario da operacao", "favorecido da operacao",
            "beneficiario da transferencia", "destinatario da transferencia", "recebedor da transferencia",
        ],
    },
    "boleto": {
        "prioridade": 120,
        "rotulos": [
            "beneficiario do boleto", "sacador avalista", "beneficiario original", "cedente",
            "nome do cedente", "beneficiario da cobranca", "recebedor da cobranca",
            "beneficiario do titulo", "cedente do titulo", "sacador do titulo",
        ],
    },
    "razao_social": {
        "prioridade": 118,
        "rotulos": [
            "razao social", "razao social do beneficiario", "razao social do recebedor",
            "razao social do favorecido", "razao social do destinatario", "razao social do cedente",
            "razao social do sacador", "nome empresarial", "nome empresarial do beneficiario",
            "razao social do estabelecimento",
        ],
    },
    "pix": {
        "prioridade": 126,
        "rotulos": [
            "dados de quem recebeu", "dados do recebedor", "dados do beneficiario", "dados do destinatario",
            "informacoes do recebedor", "informacoes do beneficiario", "recebedor da transacao",
            "beneficiario da transacao", "destinatario do pix", "favorecido do pix",
        ],
    },
    "empresarial": {
        "prioridade": 112,
        "rotulos": [
            "merchant", "merchant name", "merchant account", "merchant beneficiary", "merchant receiver",
            "estabelecimento", "nome do estabelecimento", "estabelecimento comercial", "empresa recebedora",
            "empresa beneficiaria",
        ],
    },
    "fintech": {
        "prioridade": 110,
        "rotulos": [
            "conta destinataria", "conta beneficiaria", "titular da conta destino",
            "titular da conta recebedora", "titular da conta beneficiaria", "titular da conta favorecida",
            "titular da conta do recebedor", "titular da conta pix", "titular da chave pix",
            "nome do titular da conta destino",
        ],
    },
    "cobranca": {
        "prioridade": 108,
        "rotulos": [
            "beneficiario da cobranca", "recebedor da cobranca", "favorecido da cobranca",
            "destinatario da cobranca", "empresa cobradora", "empresa recebedora da cobranca",
            "beneficiario do pagamento", "recebedor do pagamento", "destinatario do pagamento",
            "favorecido do pagamento",
        ],
    },
    "institucional": {
        "prioridade": 106,
        "rotulos": [
            "orgao recebedor", "entidade recebedora", "instituicao beneficiaria", "instituicao recebedora",
            "instituicao destinataria", "empresa destinataria", "organizacao recebedora",
            "organizacao beneficiaria", "entidade destinataria",
        ],
    },
    "alternativo": {
        "prioridade": 92,
        "rotulos": [
            "nome do titular", "titular da conta", "nome cadastrado", "nome registrado",
            "nome do cliente recebedor", "nome do cliente beneficiario", "nome da conta destino",
            "nome vinculado a chave pix", "nome vinculado a conta", "nome do portador da conta destino",
        ],
    },
}

ROTULOS_NEGATIVOS_PAGADOR = {
    "nome do pagador", "pagador", "nome do cliente", "cliente", "nome do cliente pagador", "pagante",
    "remetente", "dados do pagador", "nome do sacado", "sacado", "pagador original",
    "cliente da conta de debito", "nome da conta de debito",
}
ROTULOS_NEGATIVOS_ORIGEM = {
    "conta de debito", "conta debitada", "conta de origem", "titular da conta de origem",
    "nome do titular da conta de origem", "dados da conta de origem", "conta origem",
}
ROTULOS_NEGATIVOS_OPERACIONAIS = {
    "autenticacao", "codigo de barras", "linha digitavel", "historico", "descricao", "comprovante",
    "data do pagamento", "hora", "horario", "agencia", "nosso numero", "numero do documento",
}


def em_modo_frozen() -> bool:
    return infra_is_frozen()


def diretorio_execucao_base() -> Path:
    """Retorna base estável em modo Python normal e em executável PyInstaller."""
    return infra_app_root()


def caminho_base_usuario() -> Path:
    """Diretório base da aplicação para logs/config/temp independente do cwd."""
    return infra_data_path()


def pasta_logs_app() -> Path:
    p = caminho_base_usuario() / "logs"
    p.mkdir(parents=True, exist_ok=True)
    return p


def pasta_temp_app() -> Path:
    p = caminho_base_usuario() / "temp"
    p.mkdir(parents=True, exist_ok=True)
    return p


def diretorio_recursos_empacotados() -> Path:
    """Retorna diretório de recursos do PyInstaller (_MEIPASS) quando disponível."""
    meipass = getattr(sys, "_MEIPASS", None)
    if meipass:
        return Path(meipass)
    return diretorio_execucao_base()


def configurar_logging() -> logging.Logger:
    """Configura logger em arquivo para diagnóstico em produção (incluindo .exe)."""
    log_file = pasta_logs_app() / f"app_{datetime.now().strftime('%Y%m%d')}.log"

    logger = logging.getLogger("organizador_comprovantes")
    if logger.handlers:
        return logger

    logger.setLevel(logging.INFO)
    handler = logging.FileHandler(log_file, encoding="utf-8")
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    handler.setFormatter(formatter)
    logger.addHandler(handler)
    logger.propagate = False
    logger.info("Logging inicializado | frozen=%s | base=%s | recursos=%s", em_modo_frozen(), diretorio_execucao_base(), diretorio_recursos_empacotados())
    return logger


def formatar_diagnostico_texto(ambiente: dict[str, object] | None = None) -> str:
    amb = ambiente or verificar_ambiente_execucao()
    faltando = amb.get("faltando_modulos") or []
    avisos = amb.get("avisos") or []
    return (
        f"Aplicativo: {APP_NOME} v{APP_VERSAO}\n"
        f"Frozen (PyInstaller): {'Sim' if amb.get('frozen') else 'Não'}\n"
        f"Executável Python: {amb.get('python', '')}\n"
        f"Diretório atual: {amb.get('cwd', '')}\n"
        f"Base de execução: {amb.get('base_execucao', '')}\n"
        f"Base do usuário: {amb.get('base_usuario', '')}\n"
        f"Tesseract: {amb.get('tesseract') or 'Não encontrado'}\n"
        f"pdftoppm (Poppler): {amb.get('pdftoppm') or 'Não encontrado'}\n"
        f"Módulos ausentes: {', '.join(faltando) if faltando else 'Nenhum'}\n"
        f"Avisos: {' | '.join(avisos) if avisos else 'Nenhum'}"
    )


# =========================
# MODELOS
# =========================

@dataclass
class Comprovante:
    tipo_comprovante: str
    recebedor: str
    documento_favorecido: str
    tipo_pessoa_favorecida: str
    banco: str
    valor: str
    valor_float: float
    data_pagamento: str
    horario_pagamento: str
    origem_nome_extraido: str
    confianca_extracao: str
    confianca_label: str
    status_extracao: str
    observacoes_extracao: str
    arquivo_pdf: Path
    score_tipo: float = 0.0
    sinais_tipo: str = ""
    estrategia_leitura: str = "primeira_pagina"
    usou_ocr: bool = False


@dataclass
class FalhaProcessamento:
    arquivo: Path
    erro: str


@dataclass
class ResumoPasta:
    pasta_analisada: Path
    total_pdfs: int
    total_processados: int
    total_sem_texto: int
    total_falhas: int
    total_baixa_confianca: int
    soma_total_valores: float
    soma_valores_boletos: float
    soma_valores_pix: float
    soma_valores_transferencias: float
    soma_valores_desconhecidos: float
    quantidade_boletos: int
    quantidade_pix: int
    quantidade_transferencias: int
    quantidade_desconhecidos: int
    total_pf: int
    total_pj: int
    por_banco: dict[str, int]
    valor_por_banco: dict[str, float]
    valor_por_tipo: dict[str, float]
    quantidade_com_ocr: int = 0
    quantidade_multipagina: int = 0
    quantidade_boletos_sem_valor: int = 0
    quantidade_boletos_com_valor_valido: int = 0
    soma_boletos_potenciais: float = 0.0


# =========================
# NORMALIZAÇÃO
# =========================

def limpar_texto(texto: str) -> str:
    txt = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt


def normalizar_texto_regex(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    linhas = [re.sub(r"\s+", " ", ln).strip() for ln in texto.splitlines()]
    return "\n".join(linhas)


def limpar_nome_arquivo(nome: str) -> str:
    nome = limpar_texto(nome)
    nome = re.sub(r'[<>:"/\\|?*]', "_", nome)
    nome = re.sub(r"\s+", " ", nome).strip(" ._")
    return nome[:180] or "COMPROVANTE"


def normalizar_data(data_txt: str) -> str:
    formatos = ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"]
    data_txt = data_txt.strip()
    for fmt in formatos:
        try:
            return datetime.strptime(data_txt, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return data_txt.replace("/", "-")


def normalizar_valor(valor_txt: str) -> str:
    v = valor_txt.strip().replace("R$", "").replace(" ", "")
    if not v:
        return "NAO_IDENTIFICADO"

    try:
        if "," in v and "." in v:
            # Formato BR: 1.234,56
            v = v.replace(".", "").replace(",", ".")
        elif "," in v:
            # Formato com vírgula decimal: 560,60
            v = v.replace(",", ".")
        else:
            # Apenas ponto (decimal já normalizado) ou inteiro
            v = v
        return f"{float(v):.2f}"
    except ValueError:
        return "NAO_IDENTIFICADO"



def dividir_total_por_100(valor: float) -> float:
    return round(valor / 100.0, 2)


def formatar_moeda_br(valor: float) -> str:
    negativo = valor < 0
    valor = abs(valor)

    # Arredonda no nível de centavos para evitar casos como 1.999 -> 1,100
    total_centavos = int(round(valor * 100))
    inteiro, centavos = divmod(total_centavos, 100)

    inteiro_fmt = f"{inteiro:,}".replace(",", ".")
    out = f"R$ {inteiro_fmt},{centavos:02d}"
    return f"-{out}" if negativo else out

def valor_para_float(valor_txt: str) -> float:
    try:
        return float(normalizar_valor(valor_txt))
    except Exception:
        return 0.0


def parse_moeda_brasileira(valor_txt: str) -> float:
    """Alias semântico para manter compatibilidade na camada de renomeação."""
    return valor_para_float(valor_txt)


def somente_digitos(texto: str) -> str:
    return re.sub(r"\D", "", texto or "")


def valor_comprovante(item: Comprovante) -> float:
    if item.valor_float and item.valor_float > 0:
        return round(item.valor_float, 2)

    valor_txt = (item.valor or "").strip()
    if valor_txt:
        v = valor_para_float(valor_txt)
        if v > 0:
            return round(v, 2)

    return 0.0


# =========================
# EXTRAÇÃO (NÚCLEO)
# =========================

def extrair_primeiro(regexes: Iterable[str], texto: str, default: str = "NAO_IDENTIFICADO") -> str:
    for rx in regexes:
        m = re.search(rx, texto, flags=re.IGNORECASE | re.MULTILINE)
        if m:
            return limpar_texto(m.group(1).strip(" .:-"))
    return default


def cortar_por_rotulos_contaminantes(texto: str) -> str:
    t = texto
    for rotulo in ROTULOS_CONTAMINANTES:
        rx = re.escape(rotulo)
        t = re.split(rf"{rx}", t, maxsplit=1, flags=re.IGNORECASE)[0]
        t = re.sub(rf"{rx}.*$", "", t, flags=re.IGNORECASE)
    return t


def contem_fragmento_operacional(texto: str) -> bool:
    tu = limpar_texto(texto).upper()
    return any(r in tu for r in ROTULOS_CONTAMINANTES) or any(r in tu for r in ["VALOR", "DATA", "CODIGO", "BARRAS", "AUTENTICACAO", "HORA", "HORARIO"])


def sanitizar_candidato_nome(nome: str) -> str:
    n = limpar_texto(nome)
    n = re.split(r"\s+[|]\s+|\s{2,}", n, maxsplit=1)[0]
    n = re.sub(r"^(?:BOL|PIX|TRANF|DOC)\b\s*", "", n, flags=re.IGNORECASE)
    n = re.split(
        r"\b(?:CPF|CNPJ|CPF/CNPJ|DOCUMENTO|VALOR(?:\s+TOTAL|\s+DO\s+PAGAMENTO|\s+PAGO)?|DATA(?:\s+DO\s+(?:DEBITO|PAGAMENTO)|\s+PAGAMENTO)?|CODIGO\s+DE\s+BARRAS|LINHA\s+DIGITAVEL|INSTITUICAO\s+RECEBEDORA|BANCO\s+DESTINATARIO|AUTENTICACAO|NOME\s+BENEFICIARIO|RAZAO\s+SOCIAL\s+BENEFICIARIO|BENEFICIARIO\s+FINAL|NOME\s+DO\s+PAGADOR)\b",
        n,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    n = cortar_por_rotulos_contaminantes(n)
    n = re.sub(r"R\$\s*[\d\.,]+", "", n, flags=re.IGNORECASE)
    n = re.sub(r"\b\d{1,2}[\/\.-]\d{1,2}[\/\.-]\d{2,4}\b", "", n)
    n = re.sub(r"\b\d{4}-\d{2}-\d{2}\b", "", n)
    n = re.sub(r"\b\d{1,2}\s+\d{1,2}\s+\d{4}\b", "", n)
    n = re.sub(r"\s+", " ", n).strip(" .:-")
    return limpar_texto(n)


def extrair_bloco(texto: str, ancora_regex: str, limite: int = 500) -> str:
    m = re.search(ancora_regex + rf"[\s\S]{{0,{limite}}}", texto, flags=re.IGNORECASE)
    return m.group(0) if m else ""


def classificar_tipo_comprovante(texto: str) -> tuple[str, float, list[str]]:
    t = normalizar_texto_regex(texto).lower()

    sinais_fortes_boleto = []
    for termo in ["linha digitavel", "codigo de barras", "pagamento de boleto"]:
        if termo in t:
            sinais_fortes_boleto.append(f"boleto:{termo}")
    if sinais_fortes_boleto:
        return "boleto", 0.99, sinais_fortes_boleto

    if "beneficiario final" in t and re.search(r"\bvalor\b", t):
        return "boleto", 0.96, ["boleto:beneficiario final", "boleto:valor_contextual"]
    if "cedente" in t and "linha digitavel" in t:
        return "boleto", 0.98, ["boleto:cedente", "boleto:linha digitavel"]

    sinais_fortes_pix = []
    for termo in ["comprovante de pagamento pix", "dados de quem recebeu", "chave pix"]:
        if termo in t:
            sinais_fortes_pix.append(f"pix:{termo}")
    if len(sinais_fortes_pix) >= 2:
        return "pix", 0.99, sinais_fortes_pix

    sinais_fortes_transf = []
    for termo in ["conta favorecida", "titular da conta destino", "banco destino"]:
        if termo in t:
            sinais_fortes_transf.append(f"transferencia:{termo}")
    if len(sinais_fortes_transf) >= 2:
        return "transferencia", 0.95, sinais_fortes_transf

    sinais = {
        "boleto": {
            "linha digitavel": 5,
            "codigo de barras": 5,
            "boleto": 4,
            "pagamento de boleto": 5,
            "beneficiario final": 4,
            "nome beneficiario": 3,
            "razao social beneficiario": 4,
            "cedente": 3,
            "sacador avalista": 3,
            "beneficiario do titulo": 4,
            "titulo": 2,
        },
        "pix": {
            "pix": 5,
            "dados de quem recebeu": 4,
            "chave pix": 4,
            "chave": 2,
            "identificacao e2e": 4,
            "dados da transferencia": 3,
            "comprovante pix": 4,
        },
        "transferencia": {
            "transferencia": 5,
            "ted": 4,
            "doc": 3,
            "conta favorecida": 4,
            "titular da conta destino": 4,
            "banco destino": 3,
            "dados da transferencia": 3,
        },
    }

    score_map: dict[str, int] = {"boleto": 0, "pix": 0, "transferencia": 0}
    sinais_detectados: list[str] = []
    for tipo, mapa in sinais.items():
        for termo, peso in mapa.items():
            if re.search(rf"\b{re.escape(termo)}\b", t):
                score_map[tipo] += peso
                sinais_detectados.append(f"{tipo}:{termo}")

    tipo = max(score_map, key=score_map.get)
    maior = score_map[tipo]
    if maior <= 0:
        return "desconhecido", 0.0, []

    total = sum(score_map.values()) or 1
    score = round(maior / total, 2)
    return tipo, score, sinais_detectados


def detectar_tipo_comprovante(texto: str) -> str:
    tipo, _, _ = classificar_tipo_comprovante(texto)
    return tipo


def detectar_banco_layout(texto: str) -> str:
    t = limpar_texto(texto).upper()
    if "BRADESCO" in t:
        return "BRADESCO"
    if "ITAU" in t or "ITAU UNIBANCO" in t:
        return "ITAU"
    if "SANTANDER" in t:
        return "SANTANDER"
    if "CAIXA" in t or "CAIXA ECONOMICA" in t:
        return "CAIXA"
    if "BANCO DO BRASIL" in t or "BCO DO BRASIL" in t:
        return "BB"
    if "NUBANK" in t:
        return "NUBANK"
    if "BANCO INTER" in t or "INTER" in t:
        return "INTER"
    if "C6 BANK" in t or re.search(r"\bC6\b", t):
        return "C6"
    if "ORIGINAL" in t:
        return "ORIGINAL"
    if "NEXT" in t:
        return "NEXT"
    if "BTG" in t:
        return "BTG"
    if "SAFRA" in t:
        return "SAFRA"
    if "BANRISUL" in t:
        return "BANRISUL"
    if "SICOOB" in t:
        return "SICOOB"
    if "SICREDI" in t:
        return "SICREDI"
    return "NAO_IDENTIFICADO"


def extrair_descricao(texto: str) -> str:
    return extrair_primeiro(
        [
            r"descricao\s*[:\-]\s*(.+)",
            r"descricao\s+do\s+pagamento\s*[:\-]\s*(.+)",
            r"historico\s*[:\-]\s*(.+)",
            r"referencia\s*[:\-]\s*(.+)",
            r"descricao\s+do\s+titulo\s*[:\-]\s*(.+)",
            r"referencia\s+do\s+pagamento\s*[:\-]\s*(.+)",
        ],
        texto,
        default="",
    )


def descricao_util(descricao: str, bloqueados: set[str]) -> bool:
    d = sanitizar_candidato_nome(descricao)
    du = d.upper()
    if not d or len(d) < 3:
        return False
    if len(d.split()) < 2:
        return False
    ok, _ = validar_candidato_nome(d, bloqueados)
    if not ok:
        return False
    if any(tok in du for tok in ["CARTAO", "CONSIGNADO", "COBRANCA", "CONVENIO", "MENSALIDADE", "FINANCIAMENTO"]):
        if not re.search(r"\b(?:LTDA|S/A|SA|EIRELI|MEI|ASSOCIACAO|CLINICA|UNIVERSIDADE|EMPRESA|BANCO|SERVICOS)\b", du):
            return False
    return True


def normalizar_rotulo_semantico(rotulo: str) -> str:
    r = limpar_texto(rotulo).lower()
    r = re.sub(r"\s+", " ", r)
    return r.strip(" .:-|")


def classificar_rotulo_semantico(rotulo: str) -> tuple[str, str, int]:
    rn = normalizar_rotulo_semantico(rotulo)
    for grupo, cfg in ROTULOS_SEMANTICOS_POSITIVOS.items():
        prioridade = int(cfg.get("prioridade", 0))
        for base in cfg.get("rotulos", []):
            rb = normalizar_rotulo_semantico(str(base))
            if rn == rb or rn.startswith(rb + " ") or rb in rn:
                return "positivo", grupo, prioridade
    if rn in ROTULOS_NEGATIVOS_PAGADOR:
        return "negativo", "pagador", 0
    if rn in ROTULOS_NEGATIVOS_ORIGEM:
        return "negativo", "origem", 0
    if rn in ROTULOS_NEGATIVOS_OPERACIONAIS:
        return "negativo", "operacional", 0
    if re.search(r"\b(pagador|pagante|remetente|origem|debito|debitada|sacado|cliente)\b", rn):
        return "negativo", "pagador_origem", 0
    return "neutro", "desconhecido", 30


def slug_rotulo(rotulo: str) -> str:
    s = normalizar_rotulo_semantico(rotulo)
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")[:64] or "rotulo"


def extrair_campos_rotulados_semanticos(texto: str) -> tuple[list[tuple[str, str, int, str, str]], set[str], list[str]]:
    texto_base = normalizar_texto_regex(texto)
    linhas = texto_base.splitlines()
    candidatos: list[tuple[str, str, int, str, str]] = []
    bloqueados: set[str] = set()
    descartes: list[str] = []

    for i, linha in enumerate(linhas):
        ln = limpar_texto(linha)
        if not ln or ":" not in ln:
            continue
        rotulo, valor = ln.split(":", 1)
        rotulo = limpar_texto(rotulo)
        valor = limpar_texto(valor)
        if not valor and i + 1 < len(linhas):
            prox = limpar_texto(linhas[i + 1])
            if prox and ":" not in prox:
                valor = prox
        if "|" in valor:
            valor = limpar_texto(valor.split("|", 1)[0])
        if not valor:
            continue

        tipo, grupo, prioridade = classificar_rotulo_semantico(rotulo)
        valor_limpo = limpar_candidato_beneficiario(valor)
        if tipo == "negativo":
            if valor_limpo:
                bloqueados.add(valor_limpo.upper())
                descartes.append(f"rotulo_negativo:{rotulo}->{grupo}")
            continue
        if tipo == "positivo":
            origem = f"sem_{grupo}_{slug_rotulo(rotulo)}"
            bonus = 4 if len(valor_limpo.split()) >= 2 else 0
            candidatos.append((valor_limpo, origem, prioridade + bonus, rotulo, grupo))

    return candidatos, bloqueados, descartes


def nomes_bloqueados(texto: str) -> set[str]:
    texto_base = limpar_texto(texto)
    bloqueios = [
        r"nome\s+do\s+pagador\s*[:\-]\s*(.+)",
        r"dados\s+da\s+conta[\s\S]{0,220}?nome\s*[:\-]\s*(.+)",
        r"conta\s+de\s+debito[\s\S]{0,220}?nome\s*[:\-]\s*(.+)",
        r"titular\s+da\s+conta\s+de\s+origem\s*[:\-]\s*(.+)",
        r"nome\s+do\s+cliente\s*[:\-]\s*(.+)",
        r"cliente\s*[:\-]\s*(.+)",
        r"conta\s+origem[\s\S]{0,220}?nome\s*[:\-]\s*(.+)",
        r"conta\s+debitada[\s\S]{0,220}?nome\s*[:\-]\s*(.+)",
        r"nome\s+do\s+cliente\s+pagador\s*[:\-]\s*(.+)",
        r"dados\s+do\s+pagador\s*[:\-]\s*(.+)",
        r"dados\s+da\s+conta\s+de\s+origem[\s\S]{0,220}?nome\s*[:\-]\s*(.+)",
        r"nome\s+do\s+titular\s+da\s+conta\s+de\s+origem\s*[:\-]\s*(.+)",
        r"nome\s+da\s+conta\s+de\s+debito\s*[:\-]\s*(.+)",
        r"pagador\s+original\s*[:\-]\s*(.+)",
        r"sacado\s*[:\-]\s*(.+)",
        r"\bnome\s*[:\-]\s*([^\n\r|]+)\s*\|\s*cpf\s*[:\-]\s*[\d\.\-/]+",
    ]
    out: set[str] = set()
    for rx in bloqueios:
        for m in re.finditer(rx, texto_base, flags=re.IGNORECASE | re.MULTILINE):
            out.add(limpar_texto(m.group(1)).upper())
    return {x for x in out if x and x != "NAO_IDENTIFICADO"}


def extrair_nome_texto_contaminado(texto: str, bloqueados: set[str]) -> str:
    candidatos: list[str] = []
    for linha in texto.splitlines():
        ln = limpar_texto(linha)
        if not ln:
            continue
        if ln.upper().startswith("DESCRICAO"):
            continue
        ln = re.sub(r"^(?:BOL|PIX|TRANF|DOC)\b\s*", "", ln, flags=re.IGNORECASE)
        ln = re.split(
            r"\b(?:NOME\s+BENEFICIARIO|RAZAO\s+SOCIAL\s+BENEFICIARIO|BANCO\s+DESTINATARIO|INSTITUICAO\s+RECEBEDORA|VALOR(?:\s+TOTAL)?|DATA(?:\s+DO\s+(?:DEBITO|PAGAMENTO)|\s+PAGAMENTO)?|CODIGO\s+DE\s+BARRAS|LINHA\s+DIGITAVEL)\b",
            ln,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0]
        ln = re.sub(r"(?:NOME\s*BENEFICIARIO|RAZAO\s*SOCIAL\s*BENEFICIARIO|BANCO\s*DESTINATARIO|INSTITUICAO\s*RECEBEDORA).*$", "", ln, flags=re.IGNORECASE)
        ln = sanitizar_candidato_nome(ln)
        if ln:
            candidatos.append(ln)
    for c in candidatos:
        if candidato_nome_valido(c, bloqueados):
            return c
    return ""


def validar_candidato_nome(nome: str, bloqueados: set[str]) -> tuple[bool, str]:
    n = sanitizar_candidato_nome(nome)
    nu = n.upper()
    if len(n) < 3:
        return False, "menos de 3 caracteres"
    if nu == "NAO_IDENTIFICADO":
        return False, "valor nao identificado"
    if nu in bloqueados:
        return False, "coincide com pagador"
    if any(nu.startswith(p + " ") or nu == p for p in PREFIXOS_TECNICOS):
        return False, "prefixo tecnico"
    if "R$" in nu or re.search(r"\b\d{1,3}(?:\.\d{3})*,\d{2}\b", n):
        return False, "contem valor monetario"
    if re.search(r"\b\d{1,2}[\/\.-]\d{1,2}[\/\.-]\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b|\b\d{1,2}\s+\d{1,2}\s+\d{4}\b", n):
        return False, "contem data"
    termos_operacionais = [
        "VALOR TOTAL", "VALOR PAGO", "VALOR DO PAGAMENTO", "DATA DO DEBITO", "DATA DO PAGAMENTO",
        "DATA PAGAMENTO", "CODIGO DE BARRAS", "LINHA DIGITAVEL", "AUTENTICACAO", "TRANSACAO",
        "COMPROVANTE", "INTERNET BANKING", "MOBILE BANKING", "AGENCIA", "CONTA", "CONTA DE DEBITO",
        "CONTA ORIGEM", "DADOS DA CONTA", "BANCO DESTINATARIO", "INSTITUICAO RECEBEDORA", "PAGAMENTO",
        "BOLETO", "NOME DO PAGADOR", "HORA", "HORARIO", "TITULAR DA CONTA DE ORIGEM", "NOME DO CLIENTE", "CLIENTE", "PAGADOR"
    ]
    if any(t in nu for t in termos_operacionais):
        return False, "contem rotulo operacional"
    if nu in {"FINAL", "BENEFICIARIO", "RAZAO SOCIAL BENEFICIARIO"}:
        return False, "candidato fraco"
    if re.fullmatch(r"[\d\W]+", n):
        return False, "apenas numeros/simbolos"
    if len(re.findall(r"\d", n)) > 5:
        return False, "excesso de digitos"
    if nu in BANCOS_CONHECIDOS:
        return False, "nome de banco sem evidencia"
    if nu.startswith("BCO "):
        return False, "sigla de banco"
    if nu.startswith("BANCO ") and len(n.split()) <= 2:
        return False, "instituicao financeira generica"
    if any(x in nu for x in ["VALOR", "DATA", "CODIGO", "BARRAS", "AUTENTICACAO", "HORA", "HORARIO"]):
        return False, "contaminacao textual"
    if contem_fragmento_operacional(n):
        return False, "fragmento operacional residual"
    if len(n.split()) == 1 and n.upper() not in {"ODONTOPREV", "UNIMED", "AMIL", "HAPVIDA", "QUALICORP", "SULAMERICA", "FLEURY", "DASA", "PORTO"}:
        return False, "uma palavra generica"
    return True, "ok"


def candidato_nome_valido(nome: str, bloqueados: set[str]) -> bool:
    ok, _ = validar_candidato_nome(nome, bloqueados)
    return ok


def filtrar_candidato_nome(nome: str, bloqueados: set[str]) -> bool:
    return candidato_nome_valido(nome, bloqueados)


def extrair_documento_favorecido(texto: str) -> str:
    return extrair_primeiro(
        [
            r"cnpj\s*(?:do\s+)?beneficiario\s*[:\-]\s*([\d\./\-\*]+)",
            r"cpf\s*(?:do\s+)?beneficiario\s*[:\-]\s*([\d\./\-\*]+)",
            r"cpf/cnpj\s+de\s+quem\s+recebeu\s*[:\-]\s*([\d\./\-\*]+)",
            r"documento\s+do\s+favorecido\s*[:\-]\s*([\d\./\-\*]+)",
            r"documento\s+do\s+destinatario\s*[:\-]\s*([\d\./\-\*]+)",
            r"(?:cpf/cnpj|cnpj|cpf)\s*(?:do\s*(?:favorecido|beneficiario|recebedor|destinatario))?\s*[:\-]\s*([\d\./\-\*]+)",
        ],
        texto,
    )


def detectar_tipo_pessoa(documento: str, nome: str) -> str:
    doc = re.sub(r"\D", "", documento)
    if len(doc) == 14:
        return "PJ"
    if len(doc) == 11:
        return "PF"
    nome_u = limpar_texto(nome).upper()
    if any(tok in f" {nome_u} " for tok in [" LTDA", " S/A", " SA ", " EIRELI", " MEI", " ASSOCIACAO"]):
        return "PJ"
    return "NAO_DEFINIDO"


def extrair_valor_robusto(texto: str) -> str:
    texto_norm = normalizar_texto_regex(texto)

    candidatos: list[str] = []

    # 1) valores com rótulos fortes
    padroes_rotulados = [
        r"valor\s+total\s*[:\-]?\s*(R\$\s*[\d\.,]+|[\d]{1,3}(?:\.[\d]{3})*,[\d]{2})",
        r"valor\s+do\s+pagamento\s*[:\-]?\s*(R\$\s*[\d\.,]+|[\d]{1,3}(?:\.[\d]{3})*,[\d]{2})",
        r"valor\s+pago\s*[:\-]?\s*(R\$\s*[\d\.,]+|[\d]{1,3}(?:\.[\d]{3})*,[\d]{2})",
        r"valor\s*[:\-]?\s*(R\$\s*[\d\.,]+|[\d]{1,3}(?:\.[\d]{3})*,[\d]{2})",
    ]
    for rx in padroes_rotulados:
        candidatos.extend(m.group(1) for m in re.finditer(rx, texto_norm, flags=re.IGNORECASE | re.MULTILINE))

    # 2) valores próximos de palavras de contexto
    if not candidatos:
        padroes_contextuais = [
            r"(?:boleto|beneficiario|pagamento|liquido|titulo)[^\n\r]{0,40}?(R\$\s*[\d\.,]+|[\d]{1,3}(?:\.[\d]{3})*,[\d]{2})",
            r"(R\$\s*[\d\.,]+|[\d]{1,3}(?:\.[\d]{3})*,[\d]{2})[^\n\r]{0,40}?(?:boleto|beneficiario|pagamento|liquido|titulo)",
        ]
        for rx in padroes_contextuais:
            candidatos.extend(m.group(1) for m in re.finditer(rx, texto_norm, flags=re.IGNORECASE | re.MULTILINE))

    # 3) qualquer valor monetário explícito
    if not candidatos:
        candidatos.extend(m.group(0) for m in re.finditer(r"R\$\s*[\d\.,]+", texto_norm, flags=re.IGNORECASE))

    # 4) qualquer número monetário plausível mesmo sem R$
    if not candidatos:
        candidatos.extend(m.group(0) for m in re.finditer(r"\b[\d]{1,3}(?:\.[\d]{3})*,[\d]{2}\b", texto_norm))

    filtrados: list[str] = []
    for c in candidatos:
        linha_match = re.search(rf"[^\n\r]*{re.escape(c)}[^\n\r]*", texto_norm, flags=re.IGNORECASE)
        linha = linha_match.group(0).lower() if linha_match else ""
        if any(k in linha for k in ["tarifa", "desconto", "juros", "multa", "abatimento", "bonificacao", "encargo", "iof"]):
            continue
        filtrados.append(c)

    base = filtrados if filtrados else candidatos
    normalizados: list[tuple[float, str]] = []
    for v in base:
        nv = normalizar_valor(v)
        if nv != "NAO_IDENTIFICADO":
            try:
                fv = float(nv)
                if fv > 0:
                    normalizados.append((fv, nv))
            except Exception:
                pass

    if not normalizados:
        return "NAO_IDENTIFICADO"

    normalizados.sort(key=lambda x: x[0], reverse=True)
    return normalizados[0][1]


def extrair_data_hora_valor(texto: str) -> tuple[str, str, str]:
    valor = extrair_valor_robusto(texto)
    data = extrair_primeiro(
        [
            r"(?:data\s+do\s+debito|data\s+de\s+pagamento|data\s+pagamento|pago\s+em|data\s+e\s+hora)\s*[:\-]\s*(\d{2}[\/\.-]\d{2}[\/\.-]\d{2,4})",
            r"(\d{2}[\/\.-]\d{2}[\/\.-]\d{2,4})",
        ],
        texto,
    )
    hora = extrair_primeiro(
        [
            r"(?:data\s+e\s+hora)\s*[:\-]\s*\d{2}[\/\.-]\d{2}[\/\.-]\d{2,4}\s*(?:-|as)?\s*(\d{1,2}:\d{2}(?::\d{2})?)",
            r"(?:hora(?:rio)?|efetivado\s+as)\s*[:\-]\s*(\d{1,2}:\d{2}(?::\d{2})?)",
            r"(\d{1,2}:\d{2}(?::\d{2})?)",
        ],
        texto,
        default="NAO_INFORMADO",
    )
    return normalizar_data(data), hora, valor


def score_confianca(campos: dict[str, str], origem_nome: str) -> tuple[str, str]:
    pontos = 0
    origens_fortes = {
        "razao_social_beneficiario_final", "nome_beneficiario_final", "beneficiario_final",
        "razao_social_beneficiario", "nome_beneficiario", "beneficiario",
        "dados_de_quem_recebeu_nome", "conta_favorecida_titular", "titular_conta_destino"
    }
    origens_medias = {"favorecido", "destinatario", "recebedor", "fornecedor", "descricao_bradesco"}
    origens_fracas = {"descricao_contextual", "historico", "referencia", "fallback_nome_generico", "banco", "instituicao", "banco_destinatario", "instituicao_recebedora"}

    if origem_nome in origens_fortes:
        pontos += 30
    elif origem_nome in origens_medias:
        pontos += 18
    elif origem_nome.startswith("sem_direto_recebedor") or origem_nome.startswith("sem_pix"):
        pontos += 30
    elif origem_nome.startswith("sem_razao_social") or origem_nome.startswith("sem_boleto"):
        pontos += 26
    elif origem_nome.startswith("sem_empresarial") or origem_nome.startswith("sem_fintech"):
        pontos += 22
    elif origem_nome in origens_fracas or origem_nome == "nenhum_campo_confiavel":
        pontos += 5
    else:
        pontos += 10

    if campos.get("valor", "NAO_IDENTIFICADO") != "NAO_IDENTIFICADO":
        pontos += 20
    if campos.get("data_pagamento", "NAO_IDENTIFICADO") != "NAO_IDENTIFICADO":
        pontos += 15
    if campos.get("documento_favorecido", "NAO_IDENTIFICADO") != "NAO_IDENTIFICADO":
        pontos += 10
    if campos.get("tipo_comprovante", "desconhecido") != "desconhecido":
        pontos += 10
    if campos.get("horario_pagamento", "NAO_INFORMADO") != "NAO_INFORMADO":
        pontos += 5
    if campos.get("banco", "NAO_IDENTIFICADO") != "NAO_IDENTIFICADO":
        pontos += 2

    pontos = min(pontos, 100)
    if pontos >= 75:
        label = "alta"
    elif pontos >= 50:
        label = "media"
    else:
        label = "baixa"
    return f"{pontos/100:.2f}", label


def escolher_nome_por_prioridade(candidatos: list[tuple[str, str]], bloqueados: set[str]) -> tuple[str, str, list[str]]:
    descartes: list[str] = []
    for nome, origem in candidatos:
        nome_sanitizado = sanitizar_candidato_nome(nome)
        ok, motivo = validar_candidato_nome(nome_sanitizado, bloqueados)
        if ok:
            return nome_sanitizado, origem, descartes
        if nome_sanitizado:
            descartes.append(f"{origem}: {motivo}")
    return "NAO_IDENTIFICADO", "nenhum_campo_confiavel", descartes


@dataclass
class BeneficiarioMapeado:
    nome: str
    documento: str
    banco_relacionado: str
    tipo_pessoa: str
    origem_nome: str
    confianca_nome: str
    observacoes: str


def limpar_candidato_beneficiario(texto: str) -> str:
    t = limpar_texto(texto)
    t = re.sub(r"^(?:BOL|PIX|TRANF|DOC)\b\s*", "", t, flags=re.IGNORECASE)
    t = cortar_por_rotulos_contaminantes(t)
    t = re.sub(r"\b(?:CPF|CNPJ|CPF/CNPJ|DOCUMENTO)\b.*$", "", t, flags=re.IGNORECASE)
    t = re.sub(r"R\$\s*[\d\.,]+", "", t, flags=re.IGNORECASE)
    t = re.sub(r"\b\d{1,2}[\/\.-]\d{1,2}[\/\.-]\d{2,4}\b", "", t)
    t = re.sub(r"\b\d{4}-\d{2}-\d{2}\b", "", t)
    t = re.sub(r"\b\d{1,2}\s+\d{1,2}\s+\d{4}\b", "", t)
    t = re.sub(r"\s+", " ", t).strip(" .:-|")
    return limpar_texto(t)


def validar_candidato_beneficiario(nome: str, bloqueados: set[str], origem: str = "", grupo: str = "") -> tuple[bool, str]:
    n = limpar_candidato_beneficiario(nome)
    nu = n.upper()

    if not n or len(n) < 3:
        return False, "nome vazio ou curto"
    if nu == "NAO_IDENTIFICADO":
        return False, "nao identificado"
    origem_u = limpar_texto(origem).upper()
    grupo_u = limpar_texto(grupo).upper()
    if any(x in origem_u for x in ["PAGADOR", "ORIGEM", "DEBITO", "DEBITADA", "REMETENTE", "SACADO"]):
        return False, "origem semantica negativa"
    if grupo_u in {"PAGADOR", "ORIGEM", "PAGADOR_ORIGEM", "OPERACIONAL"}:
        return False, "grupo semantico negativo"
    if nu in bloqueados:
        return False, "coincide com pagador"
    toks_nome = set(nu.split())
    for blk in bloqueados:
        toks_blk = set(blk.split())
        if len(toks_nome) >= 2 and len(toks_blk) >= 2:
            inter = len(toks_nome.intersection(toks_blk))
            uniao = max(len(toks_nome.union(toks_blk)), 1)
            if inter / uniao >= 0.6:
                return False, "muito parecido com pagador"
    if nu in BANCOS_CONHECIDOS:
        return False, "instituicao financeira"
    if nu.startswith("BCO "):
        return False, "sigla bancaria"
    if nu.startswith("BANCO ") and len(n.split()) <= 3:
        return False, "banco generico"
    if contem_fragmento_operacional(n):
        return False, "contaminacao operacional"
    if re.search(r"\b\d{1,3}(?:\.\d{3})*,\d{2}\b", n):
        return False, "contem valor"
    if re.search(r"\b\d{1,2}[\/\.-]\d{1,2}[\/\.-]\d{2,4}\b", n) or re.search(r"\b\d{1,2}:\d{2}(?::\d{2})?\b", n):
        return False, "contem data/hora"
    if re.search(r"\b\d{11,14}\b", re.sub(r"\D", "", n)):
        return False, "contem documento colado"
    if len(re.findall(r"\d", n)) > 4:
        return False, "digitos em excesso"
    if re.fullmatch(r"[\W\d_]+", n):
        return False, "somente simbolos"
    if nu.startswith("HORA") or nu.startswith("HORARIO"):
        return False, "campo temporal"

    termos_operacionais = {
        "COMPROVANTE", "PAGAMENTO", "AUTENTICACAO", "CODIGO DE BARRAS", "LINHA DIGITAVEL",
        "CONTA DEBITADA", "CONTA ORIGEM", "AGENCIA", "CLIENTE", "PAGADOR",
        "BOLETO", "COBRANCA", "TRANSACAO BANCARIA",
    }
    if any(t in nu for t in termos_operacionais):
        return False, "termo operacional"

    if len(n.split()) == 1 and n.upper() not in {"ODONTOPREV", "UNIMED", "AMIL", "HAPVIDA", "QUALICORP", "SULAMERICA", "FLEURY", "DASA", "PORTO"}:
        return False, "uma palavra generica"

    return True, "ok"


def coletar_candidatos_beneficiario_boleto(texto: str) -> list[tuple[str, str, int]]:
    candidatos: list[tuple[str, str, int]] = []
    texto_base = limpar_texto(texto)

    padroes = [
        (r"^\s*razao\s+social\s+beneficiario\s+final[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "razao_social_beneficiario_final", 100),
        (r"^\s*nome\s+beneficiario\s+final[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "nome_beneficiario_final", 98),
        (r"^\s*beneficiario\s+final[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "beneficiario_final", 96),
        (r"^\s*razao\s+social\s+(?:do\s+)?beneficiario[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "razao_social_beneficiario", 94),
        (r"^\s*nome\s+(?:do\s+)?beneficiario[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "nome_beneficiario", 92),
        (r"^\s*beneficiario[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "beneficiario", 90),
        (r"^\s*favorecido[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "favorecido", 85),
        (r"^\s*recebedor[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "recebedor", 82),
        (r"^\s*destinatario[ \t]*[:\-]?[ \t]*([^\r\n|]+)$", "destinatario", 80),
        (r"cedente\s*[:\-]\s*(.+)", "cedente", 78),
        (r"sacador\s+avalista\s*[:\-]\s*(.+)", "sacador_avalista", 72),
        (r"beneficiario\s+do\s+titulo\s*[:\-]\s*(.+)", "beneficiario_titulo", 76),
        (r"favorecido\s+do\s+boleto\s*[:\-]\s*(.+)", "favorecido_boleto", 76),
        (r"nome\s+do\s+estabelecimento\s*[:\-]\s*(.+)", "nome_estabelecimento", 70),
        (r"fornecedor\s*[:\-]\s*(.+)", "fornecedor", 68),
    ]

    for rx, origem, score in padroes:
        for m in re.finditer(rx, texto_base, flags=re.IGNORECASE | re.MULTILINE):
            candidatos.append((m.group(1), origem, score))

    # Captura casos em que o valor do rótulo vem na linha seguinte (muito comum em comprovantes mobile).
    padroes_multilinha = [
        (r"razao\s+social\s+beneficiario\s+final\s*:?\s*(?:\n|\r\n?)\s*([^\n\r|]+)", "razao_social_beneficiario_final_multilinha", 101),
        (r"nome\s+beneficiario\s+final\s*:?\s*(?:\n|\r\n?)\s*([^\n\r|]+)", "nome_beneficiario_final_multilinha", 99),
        (r"razao\s+social\s+beneficiario\s*:?\s*(?:\n|\r\n?)\s*([^\n\r|]+)", "razao_social_beneficiario_multilinha", 97),
        (r"nome\s+beneficiario\s*:?\s*(?:\n|\r\n?)\s*([^\n\r|]+)", "nome_beneficiario_multilinha", 95),
    ]
    for rx, origem, score in padroes_multilinha:
        for m in re.finditer(rx, texto_base, flags=re.IGNORECASE | re.MULTILINE):
            candidatos.append((m.group(1), origem, score))

    for linha in texto_base.splitlines():
        ln = limpar_texto(linha)
        if not ln:
            continue
        if re.search(r"beneficiario|favorecido|recebedor|destinatario|cedente", ln, flags=re.IGNORECASE):
            candidatos.append((ln, "linha_contaminada", 40))

    return candidatos


def resolver_beneficiario_boleto(texto: str, banco_layout: str) -> BeneficiarioMapeado:
    bloqueados = nomes_bloqueados(texto)
    sem_candidatos, bloqueados_sem, descartes_sem = extrair_campos_rotulados_semanticos(texto)
    bloqueados.update(bloqueados_sem)
    documento = extrair_documento_favorecido(texto)
    banco_relacionado = extrair_primeiro([
        r"banco\s+destinatario\s*[:\-]\s*(.+)",
        r"instituicao\s+recebedora\s*[:\-]\s*(.+)",
        r"banco\s*[:\-]\s*(.+)",
    ], texto)

    candidatos = coletar_candidatos_beneficiario_boleto(texto)
    candidatos_sem: list[tuple[str, str, int, str, str]] = [(n, o, s, "", "legacy") for n, o, s in candidatos]
    candidatos_sem.extend(sem_candidatos)

    descricao = extrair_descricao(texto)
    if banco_layout == "BRADESCO" and descricao_util(descricao, bloqueados):
        candidatos_sem.append((descricao, "descricao_bradesco", 88, "descricao", "fallback"))

    descartes: list[str] = []
    validos: list[tuple[str, str, int, str, str]] = []

    for bruto, origem, score, rotulo_literal, grupo in candidatos_sem:
        limpo = limpar_candidato_beneficiario(bruto)
        ok, motivo = validar_candidato_beneficiario(limpo, bloqueados, origem=origem, grupo=grupo)
        if ok:
            validos.append((limpo, origem, score, rotulo_literal, grupo))
        elif limpo:
            descartes.append(f"{origem}: {motivo}")

    if not validos:
        fallback = extrair_nome_texto_contaminado(texto, bloqueados)
        if fallback:
            validos.append((fallback, "fallback_texto_contaminado", 60, "", "fallback"))

    if not validos:
        nome_final, origem_final, rotulo_final, grupo_final = "NAO_IDENTIFICADO", "nenhum_campo_confiavel", "", ""
        confianca = "baixa"
    else:
        validos.sort(key=lambda x: x[2], reverse=True)
        nome_final, origem_final, score_final, rotulo_final, grupo_final = validos[0]
        if score_final >= 90:
            confianca = "alta"
        elif score_final >= 70:
            confianca = "media"
        else:
            confianca = "baixa"

    obs = []
    if origem_final != "nenhum_campo_confiavel":
        obs.append(f"beneficiario mapeado por {origem_final}")
        if rotulo_final:
            obs.append(f"rotulo vencedor: {rotulo_final}")
        if grupo_final:
            obs.append(f"grupo semantico: {grupo_final}")
    else:
        obs.append("beneficiario nao encontrado")
    if descartes_sem:
        obs.append("bloqueios semanticos: " + "; ".join(descartes_sem[:6]))
    if descartes:
        obs.append("descartes: " + "; ".join(descartes[:5]))

    return BeneficiarioMapeado(
        nome=nome_final,
        documento=documento,
        banco_relacionado=banco_relacionado,
        tipo_pessoa=detectar_tipo_pessoa(documento, nome_final),
        origem_nome=origem_final,
        confianca_nome=confianca,
        observacoes=" | ".join(obs),
    )


def resolver_beneficiario_pix(texto: str, banco_layout: str) -> BeneficiarioMapeado:
    bloqueados = nomes_bloqueados(texto)
    sem_candidatos, bloqueados_sem, descartes_sem = extrair_campos_rotulados_semanticos(texto)
    bloqueados.update(bloqueados_sem)
    bloco = extrair_bloco(texto, r"dados\s+de\s+quem\s+recebeu")
    doc_bloco = extrair_primeiro([r"(?:cpf/cnpj|cpf|cnpj)\s*[:\-]\s*([\d\.\-/\*]+)"], bloco, default="")
    candidatos = [
        (extrair_primeiro([r"nome\s*[:\-]\s*(.+)"], bloco, default=""), "dados_de_quem_recebeu_nome", 140 if bloco else 96, "Nome", "pix"),
        (extrair_primeiro([r"instituicao\s*[:\-]\s*(.+)"], bloco, default=""), "instituicao", 68 if bloco else 90, "Instituicao", "pix"),
        (extrair_primeiro([r"^\s*instituicao\s+recebedora\s*[:\-]?\s*([^\n\r|]+)$"], texto), "instituicao_recebedora", 76, "Instituicao recebedora", "institucional"),
        (extrair_primeiro([r"^\s*favorecido[ 	]*[:\-]?\s*([^\n\r|]+)$"], texto), "favorecido", 86, "Favorecido", "direto_recebedor"),
    ]
    candidatos.extend(sem_candidatos)

    validos: list[tuple[str, str, int, str, str]] = []
    descartes: list[str] = []
    for bruto, origem, score, rotulo_literal, grupo in candidatos:
        limpo = limpar_candidato_beneficiario(bruto)
        ok, motivo = validar_candidato_beneficiario(limpo, bloqueados, origem=origem, grupo=grupo)
        if ok:
            validos.append((limpo, origem, score, rotulo_literal, grupo))
        elif limpo:
            descartes.append(f"{origem}: {motivo}")

    if banco_layout == "BRADESCO" and bloco and doc_bloco and validos:
        validos = [(n, o, sc + 6, r, g) if o == "dados_de_quem_recebeu_nome" else (n, o, sc, r, g) for n, o, sc, r, g in validos]

    if not validos:
        validos.append(("NAO_IDENTIFICADO", "nenhum_campo_confiavel", 0, "", ""))

    nome, origem, sc, rotulo, grupo = sorted(validos, key=lambda x: x[2], reverse=True)[0]
    doc = extrair_documento_favorecido(texto)
    banco = extrair_primeiro([r"instituicao\s*[:\-]\s*(.+)", r"banco\s*[:\-]\s*(.+)"], bloco + "\n" + texto)
    obs = [f"beneficiario mapeado por {origem}"]
    if rotulo:
        obs.append(f"rotulo vencedor: {rotulo}")
    if grupo:
        obs.append(f"grupo semantico: {grupo}")
    if banco_layout == "BRADESCO" and bloco:
        obs.append("regra bradesco pix: bloco dados de quem recebeu priorizado")
    if descartes_sem:
        obs.append("bloqueios semanticos: " + "; ".join(descartes_sem[:4]))
    if descartes:
        obs.append("descartes: " + "; ".join(descartes[:4]))
    return BeneficiarioMapeado(
        nome,
        doc,
        banco,
        detectar_tipo_pessoa(doc, nome),
        origem,
        "alta" if sc >= 90 else "media" if sc >= 70 else "baixa",
        " | ".join(obs),
    )


def resolver_beneficiario_transferencia(texto: str, banco_layout: str) -> BeneficiarioMapeado:
    bloqueados = nomes_bloqueados(texto)
    sem_candidatos, bloqueados_sem, descartes_sem = extrair_campos_rotulados_semanticos(texto)
    bloqueados.update(bloqueados_sem)
    candidatos = [
        (extrair_primeiro([r"conta\s+favorecida[\s\S]{0,280}?titular\s*[:\-]\s*(.+)"], texto), "conta_favorecida_titular", 95, "Titular conta favorecida", "fintech"),
        (extrair_primeiro([r"titular\s+da\s+conta\s+destino\s*[:\-]\s*(.+)"], texto), "titular_conta_destino", 94, "Titular da conta destino", "fintech"),
        (extrair_primeiro([r"^\s*favorecido[ 	]*[:\-]?\s*([^\n\r|]+)$"], texto), "favorecido", 82, "Favorecido", "direto_recebedor"),
    ]
    candidatos.extend(sem_candidatos)
    nome, origem = "NAO_IDENTIFICADO", "nenhum_campo_confiavel"
    grupo_final, rotulo_final = "", ""
    for bruto, org, _, rotulo, grupo in sorted(candidatos, key=lambda x: x[2], reverse=True):
        limpo = limpar_candidato_beneficiario(bruto)
        ok, _ = validar_candidato_beneficiario(limpo, bloqueados, origem=org, grupo=grupo)
        if ok:
            nome, origem = limpo, org
            grupo_final, rotulo_final = grupo, rotulo
            break
    doc = extrair_documento_favorecido(texto)
    banco = extrair_primeiro([r"banco\s+destino\s*[:\-]\s*(.+)", r"instituicao\s+favorecida\s*[:\-]\s*(.+)", r"banco\s*[:\-]\s*(.+)"], texto)
    obs = [f"beneficiario mapeado por {origem}"]
    if rotulo_final:
        obs.append(f"rotulo vencedor: {rotulo_final}")
    if grupo_final:
        obs.append(f"grupo semantico: {grupo_final}")
    if descartes_sem:
        obs.append("bloqueios semanticos: " + "; ".join(descartes_sem[:4]))
    return BeneficiarioMapeado(nome, doc, banco, detectar_tipo_pessoa(doc, nome), origem, "media", " | ".join(obs))


def resolver_beneficiario_generico(texto: str) -> BeneficiarioMapeado:
    bloqueados = nomes_bloqueados(texto)
    sem_candidatos, bloqueados_sem, descartes_sem = extrair_campos_rotulados_semanticos(texto)
    bloqueados.update(bloqueados_sem)
    candidatos = [
        (extrair_primeiro([r"^\s*beneficiario[ 	]*[:\-]?\s*([^\n\r|]+)$"], texto), "beneficiario", 80, "Beneficiario", "direto_recebedor"),
        (extrair_primeiro([r"^\s*recebedor[ 	]*[:\-]?\s*([^\n\r|]+)$"], texto), "recebedor", 78, "Recebedor", "direto_recebedor"),
        (extrair_primeiro([r"nome\s*[:\-]\s*(.+)"], texto), "fallback_nome_generico", 60, "Nome", "fallback"),
    ]
    candidatos.extend(sem_candidatos)
    nome, origem = "NAO_IDENTIFICADO", "nenhum_campo_confiavel"
    grupo_final, rotulo_final = "", ""
    for bruto, org, _, rotulo, grupo in sorted(candidatos, key=lambda x: x[2], reverse=True):
        limpo = limpar_candidato_beneficiario(bruto)
        ok, _ = validar_candidato_beneficiario(limpo, bloqueados, origem=org, grupo=grupo)
        if ok:
            nome, origem = limpo, org
            grupo_final, rotulo_final = grupo, rotulo
            break
    doc = extrair_documento_favorecido(texto)
    banco = extrair_primeiro([r"instituicao\s*[:\-]\s*(.+)", r"banco\s*[:\-]\s*(.+)"], texto)
    obs = [f"beneficiario mapeado por {origem}"]
    if rotulo_final:
        obs.append(f"rotulo vencedor: {rotulo_final}")
    if grupo_final:
        obs.append(f"grupo semantico: {grupo_final}")
    if descartes_sem:
        obs.append("bloqueios semanticos: " + "; ".join(descartes_sem[:4]))
    return BeneficiarioMapeado(nome, doc, banco, detectar_tipo_pessoa(doc, nome), origem, "baixa", " | ".join(obs))


def escolher_entre_nome_e_instituicao(nome: str, instituicao: str) -> tuple[str, str]:
    nome_limpo = limpar_candidato_beneficiario(nome)
    inst_limpa = limpar_candidato_beneficiario(instituicao)

    if not inst_limpa:
        return (nome_limpo or "NAO_IDENTIFICADO"), "nome"
    if not nome_limpo:
        return inst_limpa, "instituicao"

    nome_score = 0
    inst_score = 0

    if len(nome_limpo.split()) >= 2:
        nome_score += 10
    if len(inst_limpa.split()) >= 2:
        inst_score += 10

    if re.search(r"\b(LTDA|S/A|SA|EIRELI|IP\s+LTDA|SERVICOS)\b", nome_limpo.upper()):
        nome_score += 20
    if re.search(r"\b(LTDA|S/A|SA|EIRELI|IP\s+LTDA|SERVICOS)\b", inst_limpa.upper()):
        inst_score += 30

    if len(inst_limpa) > len(nome_limpo):
        inst_score += 10

    if inst_score > nome_score:
        return inst_limpa, "instituicao"
    return nome_limpo, "nome"


def resolver_beneficiario_por_eliminacao(texto: str, banco_layout: str) -> BeneficiarioMapeado:
    bloqueados = nomes_bloqueados(texto)

    bloqueados_extra = set(bloqueados)
    for rx in [
        r"nome\s+do\s+pagador\s*[:\-]\s*(.+)",
        r"nome\s+do\s+cliente\s*[:\-]\s*(.+)",
        r"cliente\s*[:\-]\s*(.+)",
        r"titular\s+da\s+conta\s+de\s+origem\s*[:\-]\s*(.+)",
        r"conta\s+debitada[\s\S]{0,160}?nome\s*[:\-]\s*(.+)",
        r"dados\s+da\s+conta[\s\S]{0,160}?nome\s*[:\-]\s*(.+)",
    ]:
        for m in re.finditer(rx, texto, flags=re.IGNORECASE | re.MULTILINE):
            bloqueados_extra.add(limpar_texto(m.group(1)).upper())

    candidatos: list[tuple[str, str, int]] = []
    descartes: list[str] = []

    for linha in texto.splitlines():
        ln = limpar_texto(linha)
        if not ln:
            continue

        ln = re.sub(r"^(?:beneficiario|favorecido|recebedor|destinatario|cedente|fornecedor|nome\s+beneficiario|razao\s+social\s+beneficiario)\s*[:\-]\s*", "", ln, flags=re.IGNORECASE)
        ln = sanitizar_candidato_nome(ln)
        if len(ln) < 3:
            continue

        score = 0
        lu = ln.upper()

        if len(ln.split()) >= 2:
            score += 20
        if re.search(r"\b(LTDA|S/A|SA|EIRELI|MEI|SERVICOS|HOLDING|CLINICA|ASSOCIACAO|INSTITUTO)\b", lu):
            score += 35
        if re.search(r"beneficiario|favorecido|recebedor|cedente", texto, flags=re.IGNORECASE):
            score += 10

        ok, motivo = validar_candidato_beneficiario(ln, bloqueados_extra)
        if ok:
            candidatos.append((ln, "eliminacao_geral", score))
        else:
            descartes.append(f"{ln[:60]}: {motivo}")

    instituicoes: list[tuple[str, str, int]] = []
    for rx in [
        r"instituicao\s+recebedora\s*[:\-]\s*(.+)",
        r"instituicao\s*[:\-]\s*(.+)",
    ]:
        for m in re.finditer(rx, texto, flags=re.IGNORECASE | re.MULTILINE):
            inst = limpar_candidato_beneficiario(m.group(1))
            ok, motivo = validar_candidato_beneficiario(inst, bloqueados)
            if ok:
                bonus = 30 if re.search(r"\b(LTDA|S/A|SA|EIRELI|IP\s+LTDA)\b", inst.upper()) else 10
                instituicoes.append((inst, "instituicao_reforco", bonus))
            elif inst:
                descartes.append(f"instituicao:{inst[:50]}: {motivo}")

    candidatos.extend(instituicoes)

    if not candidatos:
        return BeneficiarioMapeado(
            nome="NAO_IDENTIFICADO",
            documento=extrair_documento_favorecido(texto),
            banco_relacionado=detectar_banco_layout(texto),
            tipo_pessoa="NAO_DEFINIDO",
            origem_nome="nenhum_campo_confiavel",
            confianca_nome="baixa",
            observacoes="beneficiario nao encontrado por eliminacao",
        )

    candidatos.sort(key=lambda x: x[2], reverse=True)
    nome_final, origem_final, score_final = candidatos[0]

    instituicao_top = instituicoes[0][0] if instituicoes else ""
    nome_promovido, origem_escolha = escolher_entre_nome_e_instituicao(nome_final, instituicao_top)
    promovido_instituicao = origem_escolha == "instituicao" and nome_promovido != nome_final

    documento = extrair_documento_favorecido(texto)
    if score_final >= 45:
        conf = "alta"
    elif score_final >= 25:
        conf = "media"
    else:
        conf = "baixa"
    if promovido_instituicao and conf == "baixa":
        conf = "media"

    obs_parts = ["beneficiario definido por eliminacao geral a partir do comprovante"]
    if promovido_instituicao:
        obs_parts.append("beneficiario promovido por instituicao")
    if descartes:
        obs_parts.append("candidatos descartados: " + "; ".join(descartes[:6]))

    return BeneficiarioMapeado(
        nome=nome_promovido,
        documento=documento,
        banco_relacionado=extrair_primeiro([
            r"banco\s+destinatario\s*[:\-]\s*(.+)",
            r"instituicao\s+recebedora\s*[:\-]\s*(.+)",
            r"banco\s*[:\-]\s*(.+)",
        ], texto),
        tipo_pessoa=detectar_tipo_pessoa(documento, nome_promovido),
        origem_nome="instituicao_reforco" if promovido_instituicao else origem_final,
        confianca_nome=conf,
        observacoes=" | ".join(obs_parts),
    )


def coletar_valores_monetarios(texto: str) -> list[tuple[float, str, str]]:
    tn = normalizar_texto_regex(texto)
    achados: list[tuple[float, str, str]] = []

    for m in re.finditer(r"R\$\s*[\d\.,]+|\b\d{1,3}(?:\.\d{3})*,\d{2}\b", tn, flags=re.IGNORECASE):
        bruto = m.group(0)
        nv = normalizar_valor(bruto)
        if nv == "NAO_IDENTIFICADO":
            continue
        try:
            valor = float(nv)
        except Exception:
            continue

        inicio = max(0, m.start() - 80)
        fim = min(len(tn), m.end() + 80)
        contexto = tn[inicio:fim]
        linha_match = re.search(rf"[^\n\r]*{re.escape(bruto)}[^\n\r]*", tn, flags=re.IGNORECASE)
        linha = linha_match.group(0) if linha_match else contexto

        if valor > 0:
            achados.append((valor, bruto, linha))

    return achados


def score_valor_boleto(valor: float, linha: str) -> int:
    lu = limpar_texto(linha).upper()
    score = 0

    if any(x in lu for x in [
        "VALOR TOTAL", "VALOR DO PAGAMENTO", "VALOR PAGO", "VALOR DO TITULO", "VALOR PRINCIPAL",
        "VALOR DO DOCUMENTO", "VALOR COBRADO", "VALOR NOMINAL", "VALOR LIQUIDADO", "TOTAL A PAGAR",
        "VALOR DO BOLETO", "VALOR PRINCIPAL DO TITULO", "VALOR DA COBRANCA",
    ]):
        score += 100

    if any(x in lu for x in [
        "BOLETO", "PAGAMENTO", "BENEFICIARIO", "BENEFICIARIO FINAL", "CEDENTE", "TITULO", "LINHA DIGITAVEL", "CODIGO DE BARRAS"
    ]):
        score += 40

    if any(x in lu for x in [
        "TARIFA", "JUROS", "MULTA", "DESCONTO", "IOF", "ABATIMENTO", "BONIFICACAO", "ENCARGOS"
    ]):
        score -= 120

    if valor > 0:
        score += 10

    return score


def _normalizar_sequencia_numerica(texto: str) -> str:
    return re.sub(r"\D", "", texto or "")


def _linha_digitavel_para_barras(linha47: str) -> str:
    d = _normalizar_sequencia_numerica(linha47)
    if len(d) != 47:
        return ""
    c1 = d[0:10]
    c2 = d[10:21]
    c3 = d[21:32]
    c4 = d[32]
    c5 = d[33:47]
    livre = c1[4:9] + c2[0:10] + c3[0:10]
    barras = c1[0:4] + c4 + c5 + livre
    return barras if len(barras) == 44 else ""


def _fator_para_data_vencimento(fator: str) -> str:
    if not re.fullmatch(r"\d{4}", fator or ""):
        return "NAO_IDENTIFICADO"
    dias = int(fator)
    if dias <= 0:
        return "NAO_IDENTIFICADO"
    base = datetime(1997, 10, 7)
    dt = base + timedelta(days=dias)
    return dt.strftime("%d/%m/%Y")


def _valor_centavos_para_str(valor10: str) -> str:
    if not re.fullmatch(r"\d{10}", valor10 or ""):
        return "NAO_IDENTIFICADO"
    valor = int(valor10) / 100.0
    return f"{valor:.2f}"


def extrair_info_codigo_boleto(texto: str) -> dict[str, str]:
    tn = normalizar_texto_regex(texto)
    sequencias = re.findall(r"(?:\d[\d\.\s-]{42,70}\d)", tn)

    barras44 = ""
    linha47 = ""
    arrec48 = ""

    for bruto in sequencias:
        dig = _normalizar_sequencia_numerica(bruto)
        if len(dig) == 47 and not linha47:
            linha47 = dig
        elif len(dig) == 44 and not barras44:
            barras44 = dig
        elif len(dig) == 48 and not arrec48:
            arrec48 = dig

    if not barras44 and linha47:
        barras44 = _linha_digitavel_para_barras(linha47)

    banco_codigo = barras44[0:3] if len(barras44) == 44 else ""
    banco_nome = BANCOS_CODIGO.get(banco_codigo, "")
    valor_codigo = _valor_centavos_para_str(barras44[9:19]) if len(barras44) == 44 else "NAO_IDENTIFICADO"
    fator = barras44[5:9] if len(barras44) == 44 else ""
    vencimento = _fator_para_data_vencimento(fator) if fator else "NAO_IDENTIFICADO"

    return {
        "codigo_barras_44": barras44,
        "linha_digitavel_47": linha47,
        "linha_arrecadacao_48": arrec48,
        "banco_codigo": banco_codigo or "NAO_IDENTIFICADO",
        "banco_codigo_nome": banco_nome or "NAO_IDENTIFICADO",
        "valor_codigo": valor_codigo,
        "fator_vencimento": fator or "NAO_IDENTIFICADO",
        "data_vencimento_codigo": vencimento,
        "origem_codigo": (
            "codigo_barras_44" if barras44 else "linha_digitavel_47" if linha47 else "arrecadacao_48" if arrec48 else "nao_identificado"
        ),
    }


def _extrair_valor_por_rotulos(texto: str, rotulos: list[str], ignorar: list[str] | None = None) -> tuple[str, str]:
    tn = normalizar_texto_regex(texto)
    ign = [x.lower() for x in (ignorar or [])]
    for i, rotulo in enumerate(rotulos, start=1):
        rx = rf"{rotulo}\s*[:\-]?\s*(R\$\s*[\d\.,]+|\d{{1,3}}(?:\.\d{{3}})*,\d{{2}})"
        for m in re.finditer(rx, tn, flags=re.IGNORECASE):
            linha = re.search(rf"[^\n\r]*{re.escape(m.group(1))}[^\n\r]*", tn, flags=re.IGNORECASE)
            ltxt = (linha.group(0).lower() if linha else "")
            if any(k in ltxt for k in ign):
                continue
            nv = normalizar_valor(m.group(1))
            if nv != "NAO_IDENTIFICADO":
                return nv, f"parser_rotulo_prioridade_{i}"
    return "NAO_IDENTIFICADO", ""


def extrair_valor_boleto(texto: str) -> tuple[str, str]:
    info_codigo = extrair_info_codigo_boleto(texto)
    if info_codigo.get("valor_codigo") and info_codigo.get("valor_codigo") != "NAO_IDENTIFICADO":
        try:
            if float(info_codigo["valor_codigo"]) > 0:
                return info_codigo["valor_codigo"], f"parser_boleto_{info_codigo.get('origem_codigo','codigo')}"
        except Exception:
            pass

    rotulos_prioritarios = [
        r"valor\s+total",
        r"valor\s+do\s+pagamento",
        r"valor\s+pago",
        r"valor\s+do\s+titulo",
        r"valor\s+principal",
        r"valor\s+do\s+documento",
        r"valor\s+cobrado",
        r"valor\s+nominal",
        r"valor\s+liquidado",
        r"total\s+a\s+pagar",
        r"valor\s+do\s+boleto",
        r"valor\s+principal\s+do\s+titulo",
        r"valor\s+da\s+cobranca",
    ]
    valor, origem = _extrair_valor_por_rotulos(
        texto,
        rotulos_prioritarios,
        ["tarifa", "desconto", "juros", "multa", "iof", "abatimento", "bonificacao", "encargos"],
    )
    if valor != "NAO_IDENTIFICADO":
        return valor, f"parser_boleto_{origem}"

    candidatos = coletar_valores_monetarios(texto)
    if candidatos:
        ranqueados: list[tuple[int, float, str]] = []
        for valor_float, bruto, linha in candidatos:
            score = score_valor_boleto(valor_float, linha)
            ranqueados.append((score, valor_float, bruto))
        ranqueados.sort(key=lambda x: (x[0], x[1]), reverse=True)
        melhor_score, _, melhor_bruto = ranqueados[0]
        if melhor_score >= 0:
            nv = normalizar_valor(melhor_bruto)
            if nv != "NAO_IDENTIFICADO":
                return nv, f"parser_boleto_score_{melhor_score}"

    v = extrair_valor_robusto(texto)
    return v, "parser_boleto_maior_valor_plausivel" if v != "NAO_IDENTIFICADO" else "parser_boleto_sem_valor"


def extrair_valor_pix(texto: str) -> tuple[str, str]:
    bloco = extrair_bloco(texto, r"dados\s+da\s+transferencia", limite=900)
    universo = bloco + "\n" + texto

    rotulos = [
        r"valor\s+da\s+transferencia",
        r"valor\s+transferido",
        r"valor\s+pix",
        r"valor",
    ]

    valor, origem = _extrair_valor_por_rotulos(
        universo,
        rotulos,
        ["tarifa", "iof", "encargo", "juros", "multa"],
    )
    if valor != "NAO_IDENTIFICADO":
        return valor, f"parser_pix_{origem}"

    m = re.search(
        r"dados\s+da\s+transferencia[\s\S]{0,400}?valor\s*[:\-]?\s*(R\$\s*[\d\.,]+|\d{1,3}(?:\.\d{3})*,\d{2})",
        universo,
        flags=re.IGNORECASE,
    )
    if m:
        nv = normalizar_valor(m.group(1))
        if nv != "NAO_IDENTIFICADO":
            return nv, "parser_pix_bloco_transferencia"

    v = extrair_valor_robusto(universo)
    return v, "parser_pix_maior_valor_bloco" if v != "NAO_IDENTIFICADO" else "parser_pix_sem_valor"


def extrair_valor_transferencia(texto: str) -> tuple[str, str]:
    valor, origem = _extrair_valor_por_rotulos(texto, [r"valor\s+da\s+transferencia", r"valor\s+do\s+pagamento", r"valor"], ["tarifa","iof"])
    if valor != "NAO_IDENTIFICADO":
        return valor, f"parser_transferencia_{origem}"
    v = extrair_valor_robusto(texto)
    return v, "parser_transferencia_maior_valor_plausivel" if v != "NAO_IDENTIFICADO" else "parser_transferencia_sem_valor"


def extrair_valor_generico(texto: str) -> tuple[str, str]:
    v = extrair_valor_robusto(texto)
    return v, "parser_generico"


def _extrair_data_hora_contextual(texto: str, rotulos_data: list[str], rotulos_hora: list[str]) -> tuple[str, str]:
    data = "NAO_IDENTIFICADO"
    hora = "NAO_INFORMADO"
    for rdata in rotulos_data:
        m = re.search(rf"{rdata}\s*[:\-]?\s*(\d{{2}}[\/\.-]\d{{2}}[\/\.-]\d{{2,4}})", texto, flags=re.IGNORECASE)
        if m:
            data = normalizar_data(m.group(1)); break
    for rh in rotulos_hora:
        m = re.search(rf"{rh}\s*[:\-]?\s*(\d{{1,2}}:\d{{2}}(?::\d{{2}})?)", texto, flags=re.IGNORECASE)
        if m:
            hora = m.group(1); break
    if data == "NAO_IDENTIFICADO":
        data, hora0, _ = extrair_data_hora_valor(texto)
        if hora == "NAO_INFORMADO":
            hora = hora0
    return data, hora


def extrair_data_hora_boleto(texto: str) -> tuple[str, str]:
    return _extrair_data_hora_contextual(texto, [r"data\s+do\s+pagamento", r"data\s+do\s+debito", r"data\s+pagamento"], [r"hora", r"horario"])


def extrair_data_hora_pix(texto: str) -> tuple[str, str]:
    bloco = extrair_bloco(texto, r"dados\s+da\s+transferencia", limite=600)
    return _extrair_data_hora_contextual(bloco + "\n" + texto, [r"data\s+da\s+transferencia", r"data\s+e\s+hora", r"data"], [r"hora", r"horario"])


def extrair_data_hora_transferencia(texto: str) -> tuple[str, str]:
    return _extrair_data_hora_contextual(texto, [r"data\s+da\s+transferencia", r"efetivado\s+as", r"data\s+do\s+pagamento"], [r"efetivado\s+as", r"hora", r"horario"])


def extrair_dados_boleto(texto: str) -> dict[str, str]:
    banco_layout = detectar_banco_layout(texto)
    info_codigo = extrair_info_codigo_boleto(texto)

    beneficiario = resolver_beneficiario_boleto(texto, banco_layout)
    observacoes_beneficiario = [beneficiario.observacoes]

    if beneficiario.nome == "NAO_IDENTIFICADO" or beneficiario.confianca_nome == "baixa":
        beneficiario_fallback = resolver_beneficiario_por_eliminacao(texto, banco_layout)
        if beneficiario_fallback.nome != "NAO_IDENTIFICADO":
            beneficiario = beneficiario_fallback
            observacoes_beneficiario = [beneficiario.observacoes, "fallback eliminacao ativado"]

    data, hora = extrair_data_hora_boleto(texto)
    valor, origem_valor = extrair_valor_boleto(texto)
    banco_final = beneficiario.banco_relacionado
    if banco_final in {"", "NAO_IDENTIFICADO"} and info_codigo.get("banco_codigo_nome") not in {"", "NAO_IDENTIFICADO"}:
        banco_final = info_codigo["banco_codigo_nome"]

    obs_codigo: list[str] = []
    if info_codigo.get("origem_codigo") and info_codigo.get("origem_codigo") != "nao_identificado":
        obs_codigo.append(f"codigo boleto detectado por {info_codigo['origem_codigo']}")
    if info_codigo.get("banco_codigo") not in {"", "NAO_IDENTIFICADO"}:
        obs_codigo.append(f"banco codigo: {info_codigo['banco_codigo']}")
    if info_codigo.get("data_vencimento_codigo") not in {"", "NAO_IDENTIFICADO"}:
        obs_codigo.append(f"vencimento pelo fator: {info_codigo['data_vencimento_codigo']}")

    return {
        "tipo_comprovante": "boleto",
        "fornecedor": beneficiario.nome,
        "documento_favorecido": beneficiario.documento,
        "tipo_pessoa_favorecida": beneficiario.tipo_pessoa,
        "banco": banco_final,
        "valor": valor,
        "data_pagamento": data,
        "horario_pagamento": hora,
        "origem_nome_extraido": beneficiario.origem_nome,
        "observacoes_extracao": f"{' | '.join([o for o in observacoes_beneficiario if o])} | valor obtido por {origem_valor}{' | ' + ' | '.join(obs_codigo) if obs_codigo else ''}",
    }


def extrair_dados_pix(texto: str) -> dict[str, str]:
    banco_layout = detectar_banco_layout(texto)
    beneficiario = resolver_beneficiario_pix(texto, banco_layout)
    data, hora = extrair_data_hora_pix(texto)
    valor, origem_valor = extrair_valor_pix(texto)
    return {
        "tipo_comprovante": "pix",
        "fornecedor": beneficiario.nome,
        "documento_favorecido": beneficiario.documento,
        "tipo_pessoa_favorecida": beneficiario.tipo_pessoa,
        "banco": beneficiario.banco_relacionado,
        "valor": valor,
        "data_pagamento": data,
        "horario_pagamento": hora,
        "origem_nome_extraido": beneficiario.origem_nome,
        "observacoes_extracao": f"{beneficiario.observacoes} | valor obtido por {origem_valor}",
    }

def extrair_dados_transferencia(texto: str) -> dict[str, str]:
    banco_layout = detectar_banco_layout(texto)
    beneficiario = resolver_beneficiario_transferencia(texto, banco_layout)
    data, hora = extrair_data_hora_transferencia(texto)
    valor, origem_valor = extrair_valor_transferencia(texto)
    return {
        "tipo_comprovante": "transferencia",
        "fornecedor": beneficiario.nome,
        "documento_favorecido": beneficiario.documento,
        "tipo_pessoa_favorecida": beneficiario.tipo_pessoa,
        "banco": beneficiario.banco_relacionado,
        "valor": valor,
        "data_pagamento": data,
        "horario_pagamento": hora,
        "origem_nome_extraido": beneficiario.origem_nome,
        "observacoes_extracao": f"{beneficiario.observacoes} | valor obtido por {origem_valor}",
    }

def extrair_dados_generico(texto: str) -> dict[str, str]:
    beneficiario = resolver_beneficiario_generico(texto)
    data, hora, _ = extrair_data_hora_valor(texto)
    valor, origem_valor = extrair_valor_generico(texto)
    return {
        "tipo_comprovante": "desconhecido",
        "fornecedor": beneficiario.nome,
        "documento_favorecido": beneficiario.documento,
        "tipo_pessoa_favorecida": beneficiario.tipo_pessoa,
        "banco": beneficiario.banco_relacionado,
        "valor": valor,
        "data_pagamento": data,
        "horario_pagamento": hora,
        "origem_nome_extraido": beneficiario.origem_nome,
        "observacoes_extracao": f"{beneficiario.observacoes} | valor obtido por {origem_valor}",
    }

def extrair_dados_comprovante(texto: str) -> dict[str, str]:
    txt = TextoProcessado.from_texto(texto)
    t = txt.normalizado
    tipo, score_tipo, sinais = classificar_tipo_comprovante(t)
    if tipo == "boleto":
        campos = extrair_dados_boleto(t)
    elif tipo == "pix":
        campos = extrair_dados_pix(t)
    elif tipo == "transferencia":
        campos = extrair_dados_transferencia(t)
    else:
        campos = extrair_dados_generico(t)

    confianca, label = score_confianca(campos, campos.get("origem_nome_extraido", ""))
    campos["confianca_extracao"] = confianca
    campos["confianca_label"] = label
    campos["score_tipo"] = str(score_tipo)
    campos["sinais_tipo"] = " | ".join(sinais[:12]) if sinais else ""
    obs = campos.get("observacoes_extracao", "")
    if campos.get("documento_favorecido", "NAO_IDENTIFICADO") == "NAO_IDENTIFICADO":
        obs = (obs + " | " if obs else "") + "documento do favorecido nao identificado"
    campos["observacoes_extracao"] = obs
    return campos

def valor_para_nome_arquivo(valor: str) -> str:
    v = parse_moeda_brasileira(valor)
    if v <= 0:
        return "SEM_VALOR"
    return f"{v:.2f}".replace(".", "_")


def documento_para_nome_arquivo(documento: str) -> str:
    d = somente_digitos(documento)
    if len(d) in {11, 14}:
        return d
    return "SEM_DOC"


def data_para_nome_arquivo(data: str) -> str:
    p = re.split(r"[-/]", data or "")
    if len(p) == 3 and len(p[0]) == 4:
        yyyy, mm, dd = p
        return f"{dd} {mm} {yyyy}"
    if len(p) == 3:
        dd, mm, yyyy = p
        return f"{dd} {mm} {yyyy}"
    return (data or "SEM_DATA").replace("-", " ").replace("/", " ")


def nome_util_para_arquivo(nome: str) -> tuple[bool, str]:
    n = limpar_candidato_beneficiario(nome)
    if not n or n == "NAO_IDENTIFICADO":
        return False, "nao_identificado"
    if len(n) < 4:
        return False, "muito_curto"
    nu = n.upper()
    ruins = {
        "BENEFICIARIO", "BENEFICIARIO FINAL", "PAGAMENTO", "BOLETO", "COMPROVANTE",
        "INSTITUICAO RECEBEDORA", "NAO_IDENTIFICADO", "CONTA DEBITADA", "BANCO",
    }
    if nu in ruins or nu in BANCOS_CONHECIDOS:
        return False, "generico_ou_banco"
    if re.search(r"\b(?:R\$|VALOR|DATA|HORA|PAGAMENTO|LINHA DIGITAVEL)\b", nu):
        return False, "contaminado"
    if len(re.findall(r"\d", n)) > 6:
        return False, "excesso_digitos"
    return True, "ok"


def _deduplicar_tokens_nome(nome: str) -> str:
    toks = [t for t in limpar_nome_arquivo(nome).split() if t]
    out: list[str] = []
    for t in toks:
        if not out or out[-1] != t:
            out.append(t)
    return " ".join(out)


def score_nome_arquivo_boleto(nome_principal: str, dados: dict[str, str], origem: str) -> int:
    score = 0
    ok, _ = nome_util_para_arquivo(nome_principal)
    if ok:
        score += 40
    if dados.get("documento_favorecido") and dados.get("documento_favorecido") != "NAO_IDENTIFICADO":
        score += 20
    if parse_moeda_brasileira(dados.get("valor", "")) > 0:
        score += 20
    if dados.get("data_pagamento") and dados.get("data_pagamento") != "NAO_IDENTIFICADO":
        score += 10
    if origem in {"beneficiario_final", "nome_beneficiario_final", "razao_social_beneficiario_final", "razao_social_beneficiario", "nome_beneficiario"}:
        score += 20
    elif origem.startswith("sem_direto_recebedor") or origem.startswith("sem_pix"):
        score += 22
    elif origem.startswith("sem_razao_social") or origem.startswith("sem_boleto"):
        score += 20
    elif origem.startswith("sem_empresarial") or origem.startswith("sem_fintech"):
        score += 16
    elif origem in {"favorecido", "recebedor", "destinatario", "descricao_bradesco"}:
        score += 10
    elif origem in {"documento_fallback", "nome_neutro"}:
        score -= 5
    return max(0, min(score, 100))


def _candidatos_nome_boleto(dados: dict[str, str]) -> list[tuple[str, str, int]]:
    candidatos: list[tuple[str, str, int]] = []
    campos_ordem = [
        ("beneficiario_final", 100),
        ("razao_social_beneficiario_final", 99),
        ("nome_beneficiario_final", 98),
        ("razao_social_beneficiario", 95),
        ("nome_beneficiario", 93),
        ("beneficiario", 90),
        ("favorecido", 85),
        ("recebedor", 82),
        ("destinatario", 80),
        ("instituicao_recebedora", 74),
        ("descricao_bradesco", 72),
        ("fornecedor", 70),
    ]
    for campo, base in campos_ordem:
        valor = dados.get(campo, "")
        if valor and valor != "NAO_IDENTIFICADO":
            limpo = limpar_candidato_beneficiario(valor)
            ok, _ = nome_util_para_arquivo(limpo)
            if ok:
                candidatos.append((limpo, campo, base))
    origem_extraida = dados.get("origem_nome_extraido", "")
    fornecedor = dados.get("fornecedor", "")
    if fornecedor and fornecedor != "NAO_IDENTIFICADO" and origem_extraida.startswith("sem_"):
        limpo_fornecedor = limpar_candidato_beneficiario(fornecedor)
        ok_fornecedor, _ = nome_util_para_arquivo(limpo_fornecedor)
        if ok_fornecedor:
            bonus = 96 if origem_extraida.startswith("sem_direto_recebedor") else 94
            candidatos.append((limpo_fornecedor, origem_extraida, bonus))
    doc = documento_para_nome_arquivo(dados.get("documento_favorecido", ""))
    if doc != "SEM_DOC":
        tipo_doc = "CNPJ" if len(doc) == 14 else "CPF"
        candidatos.append((f"DOC {tipo_doc} {doc}", "documento_fallback", 60))
    candidatos.append(("BOLETO SEM IDENTIFICACAO", "nome_neutro", 20))

    unicos: dict[tuple[str, str], int] = {}
    for nome, origem, sc in candidatos:
        chave = (nome, origem)
        unicos[chave] = max(unicos.get(chave, 0), sc)
    return sorted([(n, o, sc) for (n, o), sc in unicos.items()], key=lambda x: x[2], reverse=True)


def _db_seguranca_path() -> Path:
    return caminho_base_usuario() / "seguranca_boletos.sqlite3"


def inicializar_db_seguranca() -> None:
    con = sqlite3.connect(_db_seguranca_path())
    try:
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS decisoes_boleto (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome_confirmado TEXT NOT NULL,
                documento TEXT,
                banco TEXT,
                descricao TEXT,
                origem_extraida TEXT,
                candidatos_json TEXT,
                contexto_ref TEXT,
                criado_em TEXT,
                contador_uso INTEGER DEFAULT 1
            )
            """
        )
        con.execute("CREATE INDEX IF NOT EXISTS idx_dec_doc ON decisoes_boleto(documento)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_dec_banco ON decisoes_boleto(banco)")
        con.commit()
    finally:
        con.close()


def _contexto_referencia_boleto(dados: dict[str, str]) -> str:
    partes = [
        limpar_texto(dados.get("banco", "")),
        limpar_texto(dados.get("descricao_bradesco", "")),
        limpar_texto(dados.get("origem_nome_extraido", "")),
        documento_para_nome_arquivo(dados.get("documento_favorecido", "")),
    ]
    return " | ".join([p for p in partes if p])[:280]


def consultar_db_seguranca_boleto(dados: dict[str, str]) -> tuple[str, str] | None:
    inicializar_db_seguranca()
    doc = documento_para_nome_arquivo(dados.get("documento_favorecido", ""))
    banco = limpar_texto(dados.get("banco", ""))
    contexto = _contexto_referencia_boleto(dados)
    con = sqlite3.connect(_db_seguranca_path())
    try:
        if doc != "SEM_DOC":
            row = con.execute(
                "SELECT id, nome_confirmado, contador_uso FROM decisoes_boleto WHERE documento=? ORDER BY contador_uso DESC, id DESC LIMIT 1",
                (doc,),
            ).fetchone()
            if row:
                con.execute("UPDATE decisoes_boleto SET contador_uso = contador_uso + 1 WHERE id=?", (row[0],))
                con.commit()
                return limpar_candidato_beneficiario(row[1]), "db_documento"

        row = con.execute(
            "SELECT id, nome_confirmado, contexto_ref, contador_uso FROM decisoes_boleto WHERE banco=? ORDER BY contador_uso DESC, id DESC LIMIT 20",
            (banco,),
        ).fetchall()
        melhor: tuple[int, str, int] | None = None
        for rid, nome, ctx, uso in row:
            tokens = set((ctx or "").split())
            inter = len(tokens.intersection(set(contexto.split())))
            if inter >= 3:
                cand = (inter, nome, rid)
                if melhor is None or cand[0] > melhor[0]:
                    melhor = cand
        if melhor:
            con.execute("UPDATE decisoes_boleto SET contador_uso = contador_uso + 1 WHERE id=?", (melhor[2],))
            con.commit()
            return limpar_candidato_beneficiario(melhor[1]), "db_contexto"
        return None
    finally:
        con.close()


def salvar_db_seguranca_boleto(dados: dict[str, str], nome_confirmado: str, candidatos: list[tuple[str, str, int]]) -> None:
    inicializar_db_seguranca()
    con = sqlite3.connect(_db_seguranca_path())
    try:
        con.execute(
            """
            INSERT INTO decisoes_boleto
            (nome_confirmado, documento, banco, descricao, origem_extraida, candidatos_json, contexto_ref, criado_em, contador_uso)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)
            """,
            (
                nome_confirmado,
                documento_para_nome_arquivo(dados.get("documento_favorecido", "")),
                limpar_texto(dados.get("banco", "")),
                limpar_texto(dados.get("descricao_bradesco", "")),
                dados.get("origem_nome_extraido", ""),
                json.dumps(candidatos, ensure_ascii=False),
                _contexto_referencia_boleto(dados),
                datetime.now().isoformat(timespec="seconds"),
            ),
        )
        con.commit()
    finally:
        con.close()


def _boleto_inseguro(dados: dict[str, str], score_nome: int, origem_nome: str) -> bool:
    fornecedor = limpar_candidato_beneficiario(dados.get("fornecedor", ""))
    if fornecedor == "NAO_IDENTIFICADO":
        return True
    if dados.get("confianca_extracao", "") == "baixa":
        return True
    if origem_nome in {"nenhum_campo_confiavel", "nome_neutro", "documento_fallback", "fallback_texto_contaminado"}:
        return True
    if score_nome < 55:
        return True
    return False


def revisar_boleto_manual_gui(dados: dict[str, str], nome_auto: str, candidatos: list[tuple[str, str, int]], arquivo_pdf: Path | None = None) -> tuple[str, str, bool]:
    """Janela de revisão manual para boletos inseguros. Retorna (nome, origem, salvar_db)."""
    try:
        import tkinter as tk
        from tkinter import ttk
    except Exception:
        return nome_auto, "auto_sem_gui", False

    if threading.current_thread() is not threading.main_thread():
        return nome_auto, "auto_thread_secundaria", False

    escolhido = {"nome": nome_auto, "origem": "auto", "salvar": False}

    try:
        win = tk.Toplevel()
    except Exception:
        # Ambiente sem display/GUI disponível (ex.: execução headless).
        return nome_auto, "auto_sem_display", False
    win.title("Revisão assistida de boleto")
    win.geometry("860x560")
    win.minsize(760, 500)

    frm = ttk.Frame(win, padding=12)
    frm.pack(fill="both", expand=True)
    frm.columnconfigure(0, weight=3)
    frm.columnconfigure(1, weight=2)
    frm.rowconfigure(1, weight=1)

    cab = ttk.LabelFrame(frm, text="Contexto do boleto", padding=8)
    cab.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 8))
    texto_ctx = (
        f"Arquivo: {arquivo_pdf if arquivo_pdf else '-'}\n"
        f"Tipo: {dados.get('tipo_comprovante', '-')} | Banco: {dados.get('banco', '-')}\n"
        f"Valor: {dados.get('valor', '-')} | Data: {dados.get('data_pagamento', '-')}\n"
        f"Documento: {dados.get('documento_favorecido', '-')}\n"
        f"Origem extraída: {dados.get('origem_nome_extraido', '-')}\n"
        f"Nome automático: {nome_auto}"
    )
    ttk.Label(cab, text=texto_ctx, justify="left", wraplength=820).pack(anchor="w")

    lst_frame = ttk.LabelFrame(frm, text="Candidatos sugeridos", padding=8)
    lst_frame.grid(row=1, column=0, sticky="nsew", padx=(0, 6))
    lst_frame.rowconfigure(0, weight=1)
    lst_frame.columnconfigure(0, weight=1)
    lst = tk.Listbox(lst_frame)
    lst.grid(row=0, column=0, sticky="nsew")
    for nome, origem, score in candidatos:
        lst.insert("end", f"{nome}   | origem={origem} | score={score}")
    if candidatos:
        lst.selection_set(0)

    right = ttk.LabelFrame(frm, text="Decisão", padding=8)
    right.grid(row=1, column=1, sticky="nsew")
    right.columnconfigure(0, weight=1)

    var_manual = tk.StringVar(value=nome_auto)
    ttk.Label(right, text="Nome manual (opcional):").grid(row=0, column=0, sticky="w")
    ent = ttk.Entry(right, textvariable=var_manual)
    ent.grid(row=1, column=0, sticky="ew", pady=(2, 8))

    var_salvar = tk.BooleanVar(value=False)
    ttk.Checkbutton(right, text="Salvar no banco de segurança para próximos semelhantes", variable=var_salvar).grid(row=2, column=0, sticky="w", pady=(0, 8))

    def _confirmar_candidato() -> None:
        sel = lst.curselection()
        if sel:
            idx = sel[0]
            escolhido["nome"] = candidatos[idx][0]
            escolhido["origem"] = f"manual_candidato:{candidatos[idx][1]}"
            escolhido["salvar"] = bool(var_salvar.get())
        win.destroy()

    def _confirmar_manual() -> None:
        nome = limpar_candidato_beneficiario(var_manual.get())
        ok, _ = nome_util_para_arquivo(nome)
        if not ok:
            nome = "BOLETO SEM IDENTIFICACAO"
            escolhido["origem"] = "manual_neutro"
        else:
            escolhido["origem"] = "manual_digitado"
        escolhido["nome"] = nome
        escolhido["salvar"] = bool(var_salvar.get())
        win.destroy()

    def _usar_auto() -> None:
        escolhido["nome"] = nome_auto
        escolhido["origem"] = "auto_mantido"
        escolhido["salvar"] = False
        win.destroy()

    def _usar_neutro() -> None:
        escolhido["nome"] = "BOLETO SEM IDENTIFICACAO"
        escolhido["origem"] = "manual_neutro"
        escolhido["salvar"] = bool(var_salvar.get())
        win.destroy()

    botoes = ttk.Frame(right)
    botoes.grid(row=3, column=0, sticky="ew")
    for i in range(2):
        botoes.columnconfigure(i, weight=1)

    ttk.Button(botoes, text="Usar candidato", command=_confirmar_candidato).grid(row=0, column=0, sticky="ew", padx=2, pady=2)
    ttk.Button(botoes, text="Confirmar nome manual", command=_confirmar_manual).grid(row=0, column=1, sticky="ew", padx=2, pady=2)
    ttk.Button(botoes, text="Usar nome automático", command=_usar_auto).grid(row=1, column=0, sticky="ew", padx=2, pady=2)
    ttk.Button(botoes, text="Usar nome neutro", command=_usar_neutro).grid(row=1, column=1, sticky="ew", padx=2, pady=2)

    win.transient()
    win.grab_set()
    win.wait_window()
    return str(escolhido["nome"]), str(escolhido["origem"]), bool(escolhido["salvar"])


def escolher_nome_principal_boleto(dados: dict[str, str], arquivo_pdf: Path | None = None) -> tuple[str, str, int, list[tuple[str, str, int]]]:
    candidatos = _candidatos_nome_boleto(dados)

    hit_db = consultar_db_seguranca_boleto(dados)
    if hit_db:
        nome_db, origem_db = hit_db
        sc = score_nome_arquivo_boleto(nome_db, dados, origem_db)
        return nome_db, origem_db, sc, [(nome_db, origem_db, sc)] + candidatos

    nome_auto, origem_auto, _sc = candidatos[0]
    score_auto = score_nome_arquivo_boleto(nome_auto, dados, origem_auto)

    if _boleto_inseguro(dados, score_auto, origem_auto):
        nome_rev, origem_rev, salvar = revisar_boleto_manual_gui(dados, nome_auto, candidatos, arquivo_pdf=arquivo_pdf)
        score_rev = score_nome_arquivo_boleto(nome_rev, dados, origem_rev)
        if salvar:
            salvar_db_seguranca_boleto(dados, nome_rev, candidatos)
        return nome_rev, origem_rev, score_rev, candidatos

    return nome_auto, origem_auto, score_auto, candidatos


def nome_arquivo_boleto(dados: dict[str, str]) -> str:
    nome_principal, origem, score, candidatos = escolher_nome_principal_boleto(
        dados,
        arquivo_pdf=Path(dados.get("arquivo_pdf", "")) if dados.get("arquivo_pdf") else None,
    )
    valor = valor_para_nome_arquivo(dados.get("valor", ""))
    data = data_para_nome_arquivo(dados.get("data_pagamento", ""))
    doc = documento_para_nome_arquivo(dados.get("documento_favorecido", ""))

    componentes = ["BOL", nome_principal, valor, data]
    if nome_principal == "BOLETO SEM IDENTIFICACAO" and doc != "SEM_DOC":
        componentes.insert(2, f"DOC {doc}")

    nome_final = _deduplicar_tokens_nome(" ".join(componentes).upper())
    dados["origem_nome_arquivo"] = origem
    dados["score_nome_arquivo"] = str(score)
    dados["candidatos_nome_arquivo"] = json.dumps(candidatos, ensure_ascii=False)
    dados["observacoes_extracao"] = (
        (dados.get("observacoes_extracao", "") + " | " if dados.get("observacoes_extracao") else "")
        + f"nome_arquivo_boleto por {origem} (score={score})"
    )
    return limpar_nome_arquivo(nome_final)


def nome_arquivo_pix(dados: dict[str, str]) -> str:
    principal = limpar_candidato_beneficiario(dados.get("fornecedor", "NAO_IDENTIFICADO"))
    ok, _ = nome_util_para_arquivo(principal)
    if not ok:
        principal = f"DOC {documento_para_nome_arquivo(dados.get('documento_favorecido', ''))}"
    nome = f"PIX {principal} {valor_para_nome_arquivo(dados.get('valor', ''))} {data_para_nome_arquivo(dados.get('data_pagamento', ''))}"
    return limpar_nome_arquivo(_deduplicar_tokens_nome(nome.upper()))


def nome_arquivo_transferencia(dados: dict[str, str]) -> str:
    principal = limpar_candidato_beneficiario(dados.get("fornecedor", "NAO_IDENTIFICADO"))
    ok, _ = nome_util_para_arquivo(principal)
    if not ok:
        principal = f"DOC {documento_para_nome_arquivo(dados.get('documento_favorecido', ''))}"
    nome = f"TRANF {principal} {valor_para_nome_arquivo(dados.get('valor', ''))} {data_para_nome_arquivo(dados.get('data_pagamento', ''))}"
    return limpar_nome_arquivo(_deduplicar_tokens_nome(nome.upper()))


def nome_arquivo_generico(dados: dict[str, str]) -> str:
    principal = limpar_candidato_beneficiario(dados.get("fornecedor", "NAO_IDENTIFICADO"))
    ok, _ = nome_util_para_arquivo(principal)
    if not ok:
        principal = "DOCUMENTO"
    nome = f"DOC {principal} {data_para_nome_arquivo(dados.get('data_pagamento', ''))}"
    return limpar_nome_arquivo(_deduplicar_tokens_nome(nome.upper()))


def nome_arquivo_comprovante(dados: dict[str, str]) -> str:
    tipo = (dados.get("tipo_comprovante") or "desconhecido").lower()
    if tipo == "boleto":
        return nome_arquivo_boleto(dados)
    if tipo == "pix":
        return nome_arquivo_pix(dados)
    if tipo == "transferencia":
        return nome_arquivo_transferencia(dados)
    return nome_arquivo_generico(dados)


def gerar_nome_unico(caminho: Path) -> Path:
    if not caminho.exists():
        return caminho
    i = 1
    while True:
        c = caminho.with_name(f"{caminho.stem}_{i}{caminho.suffix}")
        if not c.exists():
            return c
        i += 1


def ler_texto_primeira_pagina(pdf_path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        from PyPDF2 import PdfReader  # type: ignore
    leitor = PdfReader(str(pdf_path))
    if not leitor.pages:
        return ""
    return (leitor.pages[0].extract_text() or "").strip()


def ler_texto_pdf_completo(pdf_path: Path, max_paginas: int | None = None) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        from PyPDF2 import PdfReader  # type: ignore

    leitor = PdfReader(str(pdf_path))
    limite = len(leitor.pages) if max_paginas is None else min(len(leitor.pages), max_paginas)
    partes: list[str] = []
    for pagina in leitor.pages[:limite]:
        partes.append((pagina.extract_text() or "").strip())
    return "\n".join(p for p in partes if p).strip()


def pdf_tem_texto_suficiente(texto: str) -> bool:
    return len(re.sub(r"\s+", "", texto or "")) >= 80


def _cache_key_pdf(pdf_path: Path) -> str:
    st = pdf_path.stat()
    return f"{pdf_path.resolve()}::{st.st_mtime_ns}::{st.st_size}"


def tentar_ocr_pdf(pdf_path: Path) -> str:
    try:
        from pdf2image import convert_from_path  # type: ignore
        import pytesseract  # type: ignore
    except Exception:
        return ""

    deps = detectar_dependencias_externas()
    poppler_path = APP_CFG.ocr.force_poppler_path or (str(Path(deps["pdftoppm"]).parent) if deps.get("pdftoppm") else None)

    if APP_CFG.ocr.force_tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = APP_CFG.ocr.force_tesseract_cmd
    elif deps.get("tesseract"):
        pytesseract.pytesseract.tesseract_cmd = deps["tesseract"]

    try:
        imagens = convert_from_path(
            str(pdf_path),
            dpi=APP_CFG.ocr.dpi,
            first_page=APP_CFG.ocr.first_page,
            last_page=APP_CFG.ocr.last_page,
            poppler_path=poppler_path,
        )
        textos = [(pytesseract.image_to_string(img, lang=APP_CFG.ocr.lang) or "").strip() for img in imagens]
        return "\n".join([t for t in textos if t]).strip()
    except Exception:
        return ""


def ler_texto_pdf_inteligente(pdf_path: Path) -> tuple[str, str, bool]:
    t1 = ler_texto_primeira_pagina(pdf_path)
    if pdf_tem_texto_suficiente(t1):
        return t1, "primeira_pagina", False

    t3 = ler_texto_pdf_completo(pdf_path, max_paginas=3)
    if pdf_tem_texto_suficiente(t3):
        return t3, "ate_3_paginas", False

    tall = ler_texto_pdf_completo(pdf_path, max_paginas=None)
    if pdf_tem_texto_suficiente(tall):
        return tall, "pdf_completo", False

    return tall, "possivel_escaneado", False


def ler_texto_com_fallback_ocr(pdf_path: Path) -> tuple[str, str, bool]:
    if APP_CFG.pipeline.enable_pdf_text_cache:
        key = _cache_key_pdf(pdf_path)
        if key in PDF_TEXT_CACHE:
            return PDF_TEXT_CACHE[key][1], PDF_TEXT_CACHE[key][2], PDF_TEXT_CACHE[key][3]

    texto, estrategia, usou_ocr = ler_texto_pdf_inteligente(pdf_path)
    if pdf_tem_texto_suficiente(texto):
        if APP_CFG.pipeline.enable_pdf_text_cache:
            PDF_TEXT_CACHE[_cache_key_pdf(pdf_path)] = (time.time(), texto, estrategia, usou_ocr)
        return texto, estrategia, usou_ocr
    texto_ocr = tentar_ocr_pdf(pdf_path)
    if pdf_tem_texto_suficiente(texto_ocr):
        if APP_CFG.pipeline.enable_pdf_text_cache:
            PDF_TEXT_CACHE[_cache_key_pdf(pdf_path)] = (time.time(), texto_ocr, f"{estrategia}+ocr", True)
        return texto_ocr, f"{estrategia}+ocr", True
    if APP_CFG.pipeline.enable_pdf_text_cache:
        PDF_TEXT_CACHE[_cache_key_pdf(pdf_path)] = (time.time(), texto, estrategia, usou_ocr)
    return texto, estrategia, usou_ocr


def processar_pdf_arquivo(pdf_path: Path) -> tuple[Comprovante | None, FalhaProcessamento | None, bool]:
    try:
        pipeline = ReceiptPipeline(PipelineDeps(read_pdf_text=ler_texto_com_fallback_ocr, classify_and_extract=extrair_dados_comprovante))
        run = pipeline.process_file(pdf_path)
        texto = run.payload.get("_texto_fonte", "")
        estrategia_leitura = run.payload.get("estrategia_leitura", "desconhecida")
        usou_ocr = str(run.payload.get("usou_ocr", "False")).lower() == "true"
        sem_texto = not bool(texto.strip())
        dados = run.payload
        dados.pop("_texto_fonte", None)
        dados["arquivo_pdf"] = str(pdf_path)

        if dados.get("tipo_comprovante") == "boleto" and (
            dados.get("valor") == "NAO_IDENTIFICADO"
            or valor_para_float(dados.get("valor", "")) <= 0
            or dados.get("fornecedor") == "NAO_IDENTIFICADO"
        ):
            texto_full = ler_texto_pdf_completo(pdf_path, max_paginas=None)
            if texto_full and len(texto_full) >= len(texto):
                dados2 = extrair_dados_comprovante(texto_full)

                valor1 = valor_para_float(dados.get("valor", ""))
                valor2 = valor_para_float(dados2.get("valor", ""))

                melhorou_valor = valor2 > valor1
                melhorou_nome = (
                    dados.get("fornecedor") == "NAO_IDENTIFICADO"
                    and dados2.get("fornecedor") != "NAO_IDENTIFICADO"
                )

                if melhorou_valor or melhorou_nome:
                    dados = dados2
                    dados["arquivo_pdf"] = str(pdf_path)
                    estrategia_leitura = f"{estrategia_leitura}|reprocessamento_completo"

            if dados.get("valor") == "NAO_IDENTIFICADO" or valor_para_float(dados.get("valor", "")) <= 0:
                texto_ocr = tentar_ocr_pdf(pdf_path)
                if texto_ocr:
                    dados3 = extrair_dados_comprovante(texto_ocr)
                    if valor_para_float(dados3.get("valor", "")) > valor_para_float(dados.get("valor", "")):
                        dados = dados3
                        dados["arquivo_pdf"] = str(pdf_path)
                        estrategia_leitura = f"{estrategia_leitura}|reprocessamento_ocr"
                        usou_ocr = True

        valor_float = valor_para_float(dados.get("valor", "NAO_IDENTIFICADO"))
        item = Comprovante(
            tipo_comprovante=dados["tipo_comprovante"],
            recebedor=dados["fornecedor"],
            documento_favorecido=dados["documento_favorecido"],
            tipo_pessoa_favorecida=dados["tipo_pessoa_favorecida"],
            banco=dados["banco"],
            valor=dados["valor"],
            valor_float=valor_float,
            data_pagamento=dados["data_pagamento"],
            horario_pagamento=dados["horario_pagamento"],
            origem_nome_extraido=dados["origem_nome_extraido"],
            confianca_extracao=dados["confianca_extracao"],
            confianca_label=dados["confianca_label"],
            status_extracao="ok" if not sem_texto else "sem_texto",
            observacoes_extracao=(
                (dados.get("observacoes_extracao", "") + " | " if dados.get("observacoes_extracao") else "")
                + "tempos_etapas_ms="
                + ",".join([f"{t.stage}:{t.elapsed_ms}" for t in run.timings])
            ),
            arquivo_pdf=pdf_path,
            score_tipo=float(dados.get("score_tipo", 0.0) or 0.0),
            sinais_tipo=dados.get("sinais_tipo", ""),
            estrategia_leitura=estrategia_leitura,
            usou_ocr=usou_ocr,
        )
        return item, None, sem_texto
    except Exception as exc:
        return None, FalhaProcessamento(pdf_path, str(exc)), False


def gerar_resumo(pasta: Path, comprovantes: list[Comprovante], falhas: list[FalhaProcessamento], total_pdfs: int, total_sem_texto: int) -> ResumoPasta:
    por_banco = Counter(c.banco for c in comprovantes)
    valor_por_banco: dict[str, float] = defaultdict(float)
    valor_por_tipo: dict[str, float] = defaultdict(float)

    soma_total = 0.0
    soma_boletos = 0.0
    soma_pix = 0.0
    soma_transf = 0.0
    soma_desc = 0.0
    soma_boletos_potenciais = 0.0
    boletos_com_valor = 0
    boletos_sem_valor = 0

    for c in comprovantes:
        valor_ref = valor_comprovante(c)

        if c.tipo_comprovante == "boleto":
            if valor_ref > 0:
                soma_boletos += valor_ref
                boletos_com_valor += 1
            else:
                boletos_sem_valor += 1

        sinais = (c.sinais_tipo or "").lower()
        if (
            c.tipo_comprovante == "boleto"
            or "boleto:" in sinais
            or "linha digitavel" in sinais
            or "codigo de barras" in sinais
        ) and valor_ref > 0:
            soma_boletos_potenciais += valor_ref

        if valor_ref <= 0:
            continue

        soma_total += valor_ref
        valor_por_banco[c.banco] += valor_ref
        valor_por_tipo[c.tipo_comprovante] += valor_ref

        if c.tipo_comprovante == "pix":
            soma_pix += valor_ref
        elif c.tipo_comprovante == "transferencia":
            soma_transf += valor_ref
        elif c.tipo_comprovante != "boleto":
            soma_desc += valor_ref

    return ResumoPasta(
        pasta_analisada=pasta,
        total_pdfs=total_pdfs,
        total_processados=len(comprovantes),
        total_sem_texto=total_sem_texto,
        total_falhas=len(falhas),
        total_baixa_confianca=sum(1 for c in comprovantes if c.confianca_label == "baixa"),
        soma_total_valores=round(soma_total, 2),
        soma_valores_boletos=round(soma_boletos, 2),
        soma_valores_pix=round(soma_pix, 2),
        soma_valores_transferencias=round(soma_transf, 2),
        soma_valores_desconhecidos=round(soma_desc, 2),
        quantidade_boletos=sum(1 for c in comprovantes if c.tipo_comprovante == "boleto"),
        quantidade_pix=sum(1 for c in comprovantes if c.tipo_comprovante == "pix"),
        quantidade_transferencias=sum(1 for c in comprovantes if c.tipo_comprovante == "transferencia"),
        quantidade_desconhecidos=sum(1 for c in comprovantes if c.tipo_comprovante == "desconhecido"),
        total_pf=sum(1 for c in comprovantes if c.tipo_pessoa_favorecida == "PF"),
        total_pj=sum(1 for c in comprovantes if c.tipo_pessoa_favorecida == "PJ"),
        por_banco=dict(por_banco),
        valor_por_banco=dict(valor_por_banco),
        valor_por_tipo=dict(valor_por_tipo),
        quantidade_com_ocr=sum(1 for c in comprovantes if c.usou_ocr),
        quantidade_multipagina=sum(1 for c in comprovantes if "pdf_completo" in (c.estrategia_leitura or "") or "reprocessamento_completo" in (c.estrategia_leitura or "")),
        quantidade_boletos_sem_valor=boletos_sem_valor,
        quantidade_boletos_com_valor_valido=boletos_com_valor,
        soma_boletos_potenciais=round(soma_boletos_potenciais, 2),
    )


def mapear_pasta_e_gerar_resumo(pasta_origem: Path, recursivo: bool = True) -> tuple[list[Comprovante], ResumoPasta, list[FalhaProcessamento]]:
    if not pasta_origem.exists() or not pasta_origem.is_dir():
        raise SystemExit(f"Pasta inválida: {pasta_origem}")

    arquivos = list(pasta_origem.rglob("*.pdf")) if recursivo else list(pasta_origem.glob("*.pdf"))
    comprovantes: list[Comprovante] = []
    falhas: list[FalhaProcessamento] = []
    sem_texto = 0

    for arq in sorted(arquivos):
        item, falha, st = processar_pdf_arquivo(arq)
        if item:
            comprovantes.append(item)
        if falha:
            falhas.append(falha)
        if st:
            sem_texto += 1

    resumo = gerar_resumo(pasta_origem, comprovantes, falhas, len(arquivos), sem_texto)
    return comprovantes, resumo, falhas


# =========================
# PDF LOTE
# =========================

def pagina_inicia_novo_comprovante(texto_pagina: str) -> bool:
    t = normalizar_texto_regex(texto_pagina).lower()

    sinais_fortes = [
        "comprovante de pagamento pix",
        "comprovante de transferencia",
        "pagamento de boleto",
        "linha digitavel",
        "codigo de barras",
    ]
    if any(s in t for s in sinais_fortes):
        return True

    topo = "\n".join(t.splitlines()[:8])
    if "comprovante" in topo and ("data e hora" in topo or "dados da conta" in topo):
        return True

    return False


def agrupar_paginas_em_comprovantes(leitor_pdf) -> list[list[int]]:
    grupos: list[list[int]] = []
    atual: list[int] = []
    for idx, pagina in enumerate(leitor_pdf.pages):
        texto = (pagina.extract_text() or "").strip()
        inicia = pagina_inicia_novo_comprovante(texto)
        if not atual:
            atual = [idx]
            continue
        if inicia:
            grupos.append(atual)
            atual = [idx]
        else:
            atual.append(idx)
    if atual:
        grupos.append(atual)
    return grupos


def obter_pasta_documentos() -> Path:
    return Path.home() / "Documents"


def criar_pasta_destino(banco: str, data_pagamento: str) -> Path:
    nome = limpar_nome_arquivo(f"Comprovantes de Pagamento {banco} {data_pagamento}")
    pasta = obter_pasta_documentos() / nome
    pasta.mkdir(parents=True, exist_ok=True)
    return pasta


def dividir_pdf_e_salvar(lote_pdf: Path, pasta_destino: Path) -> tuple[list[Comprovante], int]:
    try:
        from pypdf import PdfReader, PdfWriter  # type: ignore
    except Exception:
        from PyPDF2 import PdfReader, PdfWriter  # type: ignore

    leitor = PdfReader(str(lote_pdf))
    grupos = agrupar_paginas_em_comprovantes(leitor)
    out: list[Comprovante] = []
    sem_texto = 0

    for grupo in grupos:
        textos: list[str] = []
        escritor = PdfWriter()
        for idx in grupo:
            pagina = leitor.pages[idx]
            escritor.add_page(pagina)
            txt = (pagina.extract_text() or "").strip()
            if not txt:
                sem_texto += 1
            textos.append(txt)

        texto_grupo = "\n".join([t for t in textos if t]).strip()
        dados = extrair_dados_comprovante(texto_grupo)
        dados["arquivo_pdf"] = str(lote_pdf)
        dados_obs = dados.get("observacoes_extracao", "")
        dados["observacoes_extracao"] = (dados_obs + " | " if dados_obs else "") + f"leitura multipagina usada: {len(grupo)} paginas"

        destino = gerar_nome_unico(pasta_destino / f"{nome_arquivo_comprovante(dados)}.pdf")
        with destino.open("wb") as f:
            escritor.write(f)

        out.append(
            Comprovante(
                tipo_comprovante=dados["tipo_comprovante"],
                recebedor=dados["fornecedor"],
                documento_favorecido=dados["documento_favorecido"],
                tipo_pessoa_favorecida=dados["tipo_pessoa_favorecida"],
                banco=dados["banco"],
                valor=dados["valor"],
                valor_float=valor_para_float(dados["valor"]),
                data_pagamento=dados["data_pagamento"],
                horario_pagamento=dados["horario_pagamento"],
                origem_nome_extraido=dados["origem_nome_extraido"],
                confianca_extracao=dados["confianca_extracao"],
                confianca_label=dados["confianca_label"],
                status_extracao="ok" if texto_grupo else "sem_texto",
                observacoes_extracao=dados.get("observacoes_extracao", ""),
                arquivo_pdf=destino,
                score_tipo=float(dados.get("score_tipo", 0.0) or 0.0),
                sinais_tipo=dados.get("sinais_tipo", ""),
                estrategia_leitura=f"lote_multipagina_{len(grupo)}",
                usou_ocr=False,
            )
        )

    return out, sem_texto


# =========================
# EXPORTAÇÃO EXCEL
# =========================

def _autoajustar_colunas(ws) -> None:
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 90)


def exportar_excel_completo(comprovantes: list[Comprovante], resumo: ResumoPasta, falhas: list[FalhaProcessamento], destino_excel: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()

    # Aba principal
    ws = wb.active
    ws.title = "Comprovantes"
    cabecalho = [
        "tipo_comprovante", "recebedor", "documento_favorecido", "tipo_pessoa_favorecida", "banco",
        "valor", "valor_float", "data_pagamento", "horario_pagamento", "origem_nome_extraido",
        "confianca_extracao", "confianca_label", "status_extracao", "observacoes_extracao", "arquivo_pdf", "link",
        "score_tipo", "sinais_tipo", "estrategia_leitura", "usou_ocr",
    ]
    ws.append(cabecalho)
    for i in range(1, len(cabecalho) + 1):
        ws.cell(row=1, column=i).font = Font(bold=True)

    for c in comprovantes:
        ws.append([
            c.tipo_comprovante, c.recebedor, c.documento_favorecido, c.tipo_pessoa_favorecida, c.banco,
            c.valor, valor_comprovante(c), c.data_pagamento, c.horario_pagamento, c.origem_nome_extraido,
            c.confianca_extracao, c.confianca_label, c.status_extracao, c.observacoes_extracao,
            str(c.arquivo_pdf), "Abrir", c.score_tipo, c.sinais_tipo, c.estrategia_leitura, c.usou_ocr,
        ])
        r = ws.max_row
        ws.cell(row=r, column=16).hyperlink = c.arquivo_pdf.resolve().as_uri()
        ws.cell(row=r, column=16).style = "Hyperlink"
    _autoajustar_colunas(ws)

    # Aba resumo
    wr = wb.create_sheet("Resumo")
    wr.append(["metrica", "valor"])
    for row in [
        ("pasta_analisada", str(resumo.pasta_analisada)),
        ("total_pdfs", resumo.total_pdfs),
        ("total_processados", resumo.total_processados),
        ("total_sem_texto", resumo.total_sem_texto),
        ("total_falhas", resumo.total_falhas),
        ("total_baixa_confianca", resumo.total_baixa_confianca),
        ("soma_total_valores", resumo.soma_total_valores),
        ("soma_total_valores_formatado", formatar_moeda_br(resumo.soma_total_valores)),
        ("soma_valores_boletos", resumo.soma_valores_boletos),
        ("soma_valores_boletos_formatado", formatar_moeda_br(resumo.soma_valores_boletos)),
        ("soma_valores_pix", resumo.soma_valores_pix),
        ("soma_valores_pix_formatado", formatar_moeda_br(resumo.soma_valores_pix)),
        ("soma_valores_transferencias", resumo.soma_valores_transferencias),
        ("soma_valores_transferencias_formatado", formatar_moeda_br(resumo.soma_valores_transferencias)),
        ("soma_valores_desconhecidos", resumo.soma_valores_desconhecidos),
        ("soma_valores_desconhecidos_formatado", formatar_moeda_br(resumo.soma_valores_desconhecidos)),
        ("quantidade_boletos", resumo.quantidade_boletos),
        ("quantidade_pix", resumo.quantidade_pix),
        ("quantidade_transferencias", resumo.quantidade_transferencias),
        ("quantidade_desconhecidos", resumo.quantidade_desconhecidos),
        ("total_pf", resumo.total_pf),
        ("total_pj", resumo.total_pj),
        ("quantidade_com_ocr", resumo.quantidade_com_ocr),
        ("quantidade_multipagina", resumo.quantidade_multipagina),
        ("quantidade_boletos_sem_valor", resumo.quantidade_boletos_sem_valor),
        ("quantidade_boletos_com_valor_valido", resumo.quantidade_boletos_com_valor_valido),
        ("soma_boletos_potenciais", resumo.soma_boletos_potenciais),
        ("soma_boletos_potenciais_formatado", formatar_moeda_br(resumo.soma_boletos_potenciais)),
    ]:
        wr.append(list(row))
    wr["A1"].font = Font(bold=True)
    wr["B1"].font = Font(bold=True)
    _autoajustar_colunas(wr)

    # Aba baixa confiança
    wbq = wb.create_sheet("Baixa_Confianca")
    wbq.append(["arquivo_pdf", "tipo", "recebedor", "banco", "valor", "confianca", "origem_nome"])
    for c in comprovantes:
        if c.confianca_label == "baixa":
            wbq.append([str(c.arquivo_pdf), c.tipo_comprovante, c.recebedor, c.banco, c.valor, c.confianca_extracao, c.origem_nome_extraido])
    _autoajustar_colunas(wbq)

    # Aba auditoria boletos
    wab = wb.create_sheet("Auditoria_Boletos")
    wab.append(["arquivo_pdf", "recebedor", "documento_favorecido", "banco", "valor_original", "valor_float", "data_pagamento", "origem_nome_extraido", "observacoes_extracao"])
    for c in comprovantes:
        if c.tipo_comprovante == "boleto":
            wab.append([str(c.arquivo_pdf), c.recebedor, c.documento_favorecido, c.banco, c.valor, valor_comprovante(c), c.data_pagamento, c.origem_nome_extraido, c.observacoes_extracao])
    _autoajustar_colunas(wab)

    # Aba diagnostico boletos
    wd = wb.create_sheet("Diagnostico_Boletos")
    wd.append([
        "arquivo_pdf", "tipo_comprovante", "score_tipo", "sinais_tipo",
        "valor_original", "valor_float", "recebedor", "origem_nome_extraido",
        "estrategia_leitura", "usou_ocr", "observacoes_extracao",
    ])

    for c in comprovantes:
        if c.tipo_comprovante == "boleto" or "boleto" in (c.sinais_tipo or "").lower():
            wd.append([
                str(c.arquivo_pdf),
                c.tipo_comprovante,
                c.score_tipo,
                c.sinais_tipo,
                c.valor,
                valor_comprovante(c),
                c.recebedor,
                c.origem_nome_extraido,
                c.estrategia_leitura,
                c.usou_ocr,
                c.observacoes_extracao,
            ])
    _autoajustar_colunas(wd)

    # Aba falhas
    wf = wb.create_sheet("Falhas")
    wf.append(["arquivo", "erro"])
    for f in falhas:
        wf.append([str(f.arquivo), f.erro])
    _autoajustar_colunas(wf)

    destino_excel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino_excel)
    return destino_excel


def mapear_pasta_e_alimentar_excel(pasta_origem: Path, excel_saida: Path | None = None, recursivo: bool = True) -> Path:
    comprovantes, resumo, falhas = mapear_pasta_e_gerar_resumo(pasta_origem, recursivo=recursivo)
    saida = excel_saida or (pasta_origem / "relatorio_comprovantes.xlsx")
    return exportar_excel_completo(comprovantes, resumo, falhas, saida)


# =========================
# INTERFACE
# =========================

def abrir_pasta(caminho: Path) -> None:
    try:
        destino = caminho if caminho.is_dir() else caminho.parent
        if os.name == "nt":
            os.startfile(str(destino))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.run(["open", str(destino)], check=False)
        else:
            subprocess.run(["xdg-open", str(destino)], check=False)
    except Exception:
        pass


def selecionar_pasta_tk() -> Path | None:
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        pasta = filedialog.askdirectory(title="Selecione a pasta")
        root.destroy()
        return Path(pasta) if pasta else None
    except Exception:
        return None


def selecionar_arquivo_pdf_tk() -> Path | None:
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        arq = filedialog.askopenfilename(title="Selecione o PDF lote", filetypes=[("PDF", "*.pdf")])
        root.destroy()
        return Path(arq) if arq else None
    except Exception:
        return None


def selecionar_arquivo_pdf_manual() -> Path | None:
    """Seleciona PDF lote exclusivamente por interface gráfica (sem terminal)."""
    return selecionar_arquivo_pdf_tk()


def localizar_executavel(nome: str, subpastas: list[str] | None = None) -> str:
    """Localiza executável no PATH e em pastas previsíveis ao lado do app/.exe."""
    subpastas = subpastas or []
    candidatos: list[Path] = []

    achado = shutil.which(nome)
    if achado:
        candidatos.append(Path(achado))

    # Compatibilidade: se nome já vier com .exe, tenta versão sem extensão e vice-versa.
    alt = nome[:-4] if nome.lower().endswith('.exe') else f"{nome}.exe"
    achado_alt = shutil.which(alt)
    if achado_alt:
        candidatos.append(Path(achado_alt))

    base = diretorio_execucao_base()
    recursos = diretorio_recursos_empacotados()
    for raiz in (base, recursos):
        candidatos.append(raiz / nome)
        candidatos.append(raiz / alt)
        candidatos.append(raiz / "tools" / nome)
        candidatos.append(raiz / "tools" / alt)
        for sub in subpastas:
            candidatos.append(raiz / sub / nome)
            candidatos.append(raiz / sub / alt)

    for caminho in candidatos:
        if caminho.exists():
            return str(caminho.resolve())
    return ""


def detectar_dependencias_externas() -> dict[str, str]:
    return {
        "tesseract": localizar_executavel("tesseract.exe", ["tesseract", "bin"]),
        "pdftoppm": localizar_executavel("pdftoppm.exe", ["poppler", "poppler/bin", "bin"]),
    }


def verificar_ambiente_execucao() -> dict[str, object]:
    """Valida itens críticos para manter paridade entre VS Code e executável."""
    base_diag = run_environment_diagnostics()
    faltando_modulos: list[str] = list(base_diag.get("missing_python_modules", []))

    deps = detectar_dependencias_externas()
    tesseract = deps.get("tesseract", "")
    pdftoppm = deps.get("pdftoppm", "")
    avisos: list[str] = []

    if faltando_modulos:
        avisos.append(f"Módulos ausentes: {', '.join(faltando_modulos)}")
    if not tesseract:
        avisos.append("OCR pode falhar: executável 'tesseract' não encontrado no PATH")
    if not pdftoppm:
        avisos.append("Conversão OCR pode falhar: utilitário 'pdftoppm' (Poppler) não encontrado no PATH")

    return {
        "frozen": em_modo_frozen(),
        "python": sys.executable,
        "cwd": str(caminho_base_usuario()),
        "base_execucao": str(diretorio_execucao_base()),
        "base_usuario": str(caminho_base_usuario()),
        "faltando_modulos": faltando_modulos,
        "tesseract": tesseract or "",
        "pdftoppm": pdftoppm or "",
        "ocr_config": {
            "dpi": APP_CFG.ocr.dpi,
            "lang": APP_CFG.ocr.lang,
            "first_page": APP_CFG.ocr.first_page,
            "last_page": APP_CFG.ocr.last_page,
            "force_tesseract_cmd": APP_CFG.ocr.force_tesseract_cmd,
            "force_poppler_path": APP_CFG.ocr.force_poppler_path,
        },
        "writable_data_dir": bool(base_diag.get("writable_data_dir", True)),
        "avisos": avisos,
    }


def executar_dashboard(
    pasta_inicial: Path | None = None,
    lote_inicial: Path | None = None,
    auto_mapear: bool = False,
    auto_lote: bool = False,
) -> None:
    import json
    import queue
    import threading
    import time
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class ToolTip:
        def __init__(self, widget: tk.Widget, text: str):
            self.widget = widget
            self.text = text
            self.tipwindow = None
            widget.bind("<Enter>", self._show)
            widget.bind("<Leave>", self._hide)

        def _show(self, _event=None):
            if self.tipwindow:
                return
            x = self.widget.winfo_rootx() + 20
            y = self.widget.winfo_rooty() + 20
            self.tipwindow = tw = tk.Toplevel(self.widget)
            tw.wm_overrideredirect(True)
            tw.wm_geometry(f"+{x}+{y}")
            label = ttk.Label(tw, text=self.text, relief="solid", borderwidth=1, padding=4)
            label.pack()

        def _hide(self, _event=None):
            if self.tipwindow:
                self.tipwindow.destroy()
                self.tipwindow = None

    class AppDashboard:
        CONFIG_PATH = Path.home() / ".organizador_comprovantes_ui.json"

        def __init__(
            self,
            root: tk.Tk,
            pasta_inicial: Path | None = None,
            lote_inicial: Path | None = None,
            auto_mapear: bool = False,
            auto_lote: bool = False,
        ):
            self.root = root
            self.root.title(f"{APP_NOME} - Painel Principal")
            self.root.geometry("1400x860")
            self.root.minsize(1100, 700)

            # Estado
            self.dataset_bruto: list[Comprovante] = []
            self.dataset_filtrado: list[Comprovante] = []
            self.dataset_baixa: list[Comprovante] = []
            self.falhas: list[FalhaProcessamento] = []
            self.item_selecionado: Comprovante | None = None
            self.ultimo_resumo: ResumoPasta | None = None
            self.ultimo_excel: Path | None = None
            self.processando = False
            self.cancelar = threading.Event()
            self.fila_ui: queue.Queue = queue.Queue()
            self.worker: threading.Thread | None = None
            self.inicio_processamento = 0.0
            self.total_alvo = 0
            self.processados = 0
            self.coluna_ordenacao = ""
            self.ordem_desc = False
            self.revisao_manual: dict[str, dict[str, str | bool]] = {}
            self.logger = configurar_logging()
            self.pasta_inicial = pasta_inicial
            self.lote_inicial = lote_inicial
            self.auto_mapear = auto_mapear
            self.auto_lote = auto_lote

            self._estado_vars()
            self._configurar_estilo()
            self._construir_layout()
            self._carregar_config()
            self._aplicar_bindings()
            self._aplicar_estado_botoes()
            self._log("Dashboard iniciado.")
            ambiente = verificar_ambiente_execucao()
            for aviso in ambiente["avisos"]:
                self._log(f"[AVISO] {aviso}")
            if ambiente["avisos"]:
                self.var_status.set("Ambiente com avisos (ver log)")
            self._aplicar_entrada_inicial()
            self._poll_fila()

        def _configurar_estilo(self) -> None:
            style = ttk.Style(self.root)
            try:
                style.theme_use("vista")
            except Exception:
                pass
            style.configure("TLabel", font=("Segoe UI", 9))
            style.configure("TButton", font=("Segoe UI", 9), padding=(8, 5))
            style.configure("TitleDash.TLabel", font=("Segoe UI", 13, "bold"))
            style.configure("Card.TLabelframe", padding=8)

        # ---------- Estado ----------
        def _estado_vars(self) -> None:
            self.var_pasta = tk.StringVar()
            self.var_status = tk.StringVar(value="Pronto para iniciar")
            self.var_progresso_texto = tk.StringVar(value="Aguardando ação do usuário...")
            self.var_busca = tk.StringVar()
            self.var_f_tipo = tk.StringVar(value="Todos")
            self.var_f_banco = tk.StringVar(value="Todos")
            self.var_f_conf = tk.StringVar(value="Todos")
            self.var_f_status = tk.StringVar(value="Todos")
            self.var_f_pessoa = tk.StringVar(value="Todos")
            self.var_total_filtrado = tk.StringVar(value="0 / 0")
            self.var_pasta_ok = tk.StringVar(value="Pasta não selecionada")

            self.cards = {
                "pdfs": tk.StringVar(value="0"),
                "proc": tk.StringVar(value="0"),
                "valor": tk.StringVar(value="R$ 0,00"),
                "baixa": tk.StringVar(value="0"),
                "sem": tk.StringVar(value="0"),
                "falhas": tk.StringVar(value="0"),
                "%baixa": tk.StringVar(value="0%"),
                "dominante": tk.StringVar(value="-"),
                "valor_boletos": tk.StringVar(value="R$ 0,00"),
                "boletos_validos": tk.StringVar(value="0"),
            }

        def _aplicar_entrada_inicial(self) -> None:
            if self.pasta_inicial:
                self.var_pasta.set(str(self.pasta_inicial))
                self._validar_pasta_ui()
                if self.auto_mapear:
                    self.root.after(200, self._on_mapear)

            if self.lote_inicial and self.auto_lote:
                self.root.after(250, lambda: self._processar_lote_arquivo(self.lote_inicial))

        # ---------- Construção UI ----------
        def _construir_layout(self) -> None:
            self.root.columnconfigure(0, weight=1)
            self.root.rowconfigure(3, weight=1)

            self.frm_topo = ttk.Frame(self.root, padding=10)
            self.frm_topo.grid(row=0, column=0, sticky="ew")
            self.frm_topo.columnconfigure(1, weight=1)

            ttk.Label(self.frm_topo, text=f"{APP_NOME} — Painel Operacional", style="TitleDash.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0,6))
            ttk.Label(self.frm_topo, text=f"Versão {APP_VERSAO}").grid(row=0, column=2, sticky="e")
            ttk.Label(self.frm_topo, text="Pasta de trabalho:").grid(row=1, column=0, sticky="w", padx=(0, 6))
            self.ent_pasta = ttk.Entry(self.frm_topo, textvariable=self.var_pasta)
            self.ent_pasta.grid(row=1, column=1, sticky="ew")
            self.lbl_pasta_ok = ttk.Label(self.frm_topo, textvariable=self.var_pasta_ok)
            self.lbl_pasta_ok.grid(row=2, column=1, sticky="w", pady=(4, 0))

            self.frm_botoes = ttk.Frame(self.frm_topo)
            self.frm_botoes.grid(row=1, column=2, rowspan=2, padx=(8, 0), sticky="ne")

            self.btn_sel_pasta = ttk.Button(self.frm_botoes, text="Selecionar pasta", command=self._on_selecionar_pasta)
            self.btn_mapear = ttk.Button(self.frm_botoes, text="Iniciar Inventário", command=self._on_mapear)
            self.btn_lote = ttk.Button(self.frm_botoes, text="Extrair PDF lote", command=self._on_processar_lote_manual)
            self.btn_atualizar = ttk.Button(self.frm_botoes, text="Reprocessar Pasta", command=self._on_mapear)
            self.btn_exportar = ttk.Button(self.frm_botoes, text="Exportar Excel", command=self._on_exportar_excel)
            self.btn_abrir_excel = ttk.Button(self.frm_botoes, text="Abrir Excel", command=self._on_abrir_excel)
            self.btn_abrir_pasta = ttk.Button(self.frm_botoes, text="Abrir pasta", command=self._on_abrir_pasta)
            self.btn_limpar = ttk.Button(self.frm_botoes, text="Limpar Busca/Filtros", command=self._on_limpar_filtros)
            self.btn_cancelar = ttk.Button(self.frm_botoes, text="Cancelar processamento", command=self._on_cancelar)

            botoes = [
                self.btn_sel_pasta, self.btn_mapear, self.btn_lote, self.btn_atualizar, self.btn_exportar,
                self.btn_abrir_excel, self.btn_abrir_pasta, self.btn_limpar, self.btn_cancelar,
            ]
            for i, b in enumerate(botoes):
                b.grid(row=i // 2, column=i % 2, padx=2, pady=2, sticky="ew")

            tips = {
                self.btn_sel_pasta: "Seleciona a pasta principal para análise.",
                self.btn_mapear: "Inicia o inventário da pasta selecionada e atualiza os painéis.",
                self.btn_lote: "Seleciona um PDF lote e separa comprovantes por página com geração de Excel.",
                self.btn_exportar: "Exporta relatório Excel completo.",
                self.btn_cancelar: "Cancela processamento em andamento.",
            }
            for w, t in tips.items():
                ToolTip(w, t)

            self.frm_filtros = ttk.LabelFrame(self.root, text="Filtros e busca", padding=8)
            self.frm_filtros.grid(row=1, column=0, sticky="ew", padx=8, pady=(0, 8))
            for i in range(12):
                self.frm_filtros.columnconfigure(i, weight=1)

            ttk.Label(self.frm_filtros, text="Busca:").grid(row=0, column=0, sticky="w")
            self.ent_busca = ttk.Entry(self.frm_filtros, textvariable=self.var_busca)
            self.ent_busca.grid(row=0, column=1, columnspan=3, sticky="ew", padx=(4, 8))

            self.cb_tipo = self._mk_combo(self.frm_filtros, "Tipo", self.var_f_tipo, 4)
            self.cb_banco = self._mk_combo(self.frm_filtros, "Banco", self.var_f_banco, 5)
            self.cb_conf = self._mk_combo(self.frm_filtros, "Confiança", self.var_f_conf, 6)
            self.cb_status = self._mk_combo(self.frm_filtros, "Status", self.var_f_status, 7)
            self.cb_pessoa = self._mk_combo(self.frm_filtros, "PF/PJ", self.var_f_pessoa, 8)

            ttk.Label(self.frm_filtros, text="Filtrados:").grid(row=0, column=9, sticky="e")
            ttk.Label(self.frm_filtros, textvariable=self.var_total_filtrado).grid(row=0, column=10, sticky="w")

            self.frm_cards = ttk.Frame(self.root, padding=(8, 0))
            self.frm_cards.grid(row=2, column=0, sticky="ew")
            for i in range(10):
                self.frm_cards.columnconfigure(i, weight=1)

            self._mk_card(0, "PDFs encontrados", self.cards["pdfs"])
            self._mk_card(1, "Processados", self.cards["proc"])
            self._mk_card(2, "Valor total", self.cards["valor"])
            self._mk_card(3, "Baixa confiança", self.cards["baixa"])
            self._mk_card(4, "Sem texto", self.cards["sem"])
            self._mk_card(5, "Falhas", self.cards["falhas"])
            self._mk_card(6, "% baixa conf.", self.cards["%baixa"])
            self._mk_card(7, "Dominante", self.cards["dominante"])
            self._mk_card(8, "Valor boletos", self.cards["valor_boletos"])
            self._mk_card(9, "Boletos com valor válido", self.cards["boletos_validos"])

            self.notebook = ttk.Notebook(self.root)
            self.notebook.grid(row=3, column=0, sticky="nsew", padx=8, pady=8)

            self.tab_resumo = ttk.Frame(self.notebook)
            self.tab_comp = ttk.Frame(self.notebook)
            self.tab_baixa = ttk.Frame(self.notebook)
            self.tab_falhas = ttk.Frame(self.notebook)
            self.tab_revisao = ttk.Frame(self.notebook)

            self.notebook.add(self.tab_resumo, text="Resumo")
            self.notebook.add(self.tab_comp, text="Comprovantes")
            self.notebook.add(self.tab_baixa, text="Baixa Confiança")
            self.notebook.add(self.tab_falhas, text="Falhas")
            self.notebook.add(self.tab_revisao, text="Revisão Manual")

            self._build_tab_resumo()
            self._build_tab_comprovantes()
            self._build_tab_baixa()
            self._build_tab_falhas()
            self._build_tab_revisao()

            self.frm_rodape = ttk.Frame(self.root, padding=8)
            self.frm_rodape.grid(row=4, column=0, sticky="ew")
            self.frm_rodape.columnconfigure(1, weight=1)

            self.progress = ttk.Progressbar(self.frm_rodape, mode="determinate", maximum=100)
            self.progress.grid(row=0, column=0, sticky="ew", padx=(0, 8))
            self.lbl_prog = ttk.Label(self.frm_rodape, textvariable=self.var_progresso_texto)
            self.lbl_prog.grid(row=0, column=1, sticky="w")
            self.lbl_status = ttk.Label(self.frm_rodape, textvariable=self.var_status)
            self.lbl_status.grid(row=0, column=2, sticky="e")

            self.txt_log = tk.Text(self.frm_rodape, height=6, wrap="word")
            self.txt_log.grid(row=1, column=0, columnspan=3, sticky="ew", pady=(6, 0))
            self.txt_log.configure(state="disabled")

        def _mk_combo(self, parent, label, var, col):
            frm = ttk.Frame(parent)
            frm.grid(row=0, column=col, sticky="ew", padx=4)
            ttk.Label(frm, text=label).pack(anchor="w")
            cb = ttk.Combobox(frm, textvariable=var, state="readonly", values=["Todos"])
            cb.pack(fill="x")
            return cb

        def _mk_card(self, col: int, titulo: str, var: tk.StringVar) -> None:
            card = ttk.LabelFrame(self.frm_cards, text=titulo, padding=6)
            card.grid(row=0, column=col, sticky="ew", padx=3, pady=4)
            ttk.Label(card, textvariable=var, font=("Segoe UI", 11, "bold")).pack(anchor="center")

        # ---------- Abas ----------
        def _build_tab_resumo(self) -> None:
            self.tab_resumo.columnconfigure(0, weight=1)
            self.tab_resumo.rowconfigure(0, weight=1)
            self.txt_resumo = tk.Text(self.tab_resumo, wrap="word")
            self.txt_resumo.grid(row=0, column=0, sticky="nsew")
            self.txt_resumo.insert("end", "Sem análise carregada.")
            self.txt_resumo.configure(state="disabled")

        def _build_tab_comprovantes(self) -> None:
            self.tab_comp.columnconfigure(0, weight=3)
            self.tab_comp.columnconfigure(1, weight=2)
            self.tab_comp.rowconfigure(0, weight=1)

            frame_tbl = ttk.Frame(self.tab_comp)
            frame_tbl.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
            frame_tbl.columnconfigure(0, weight=1)
            frame_tbl.rowconfigure(0, weight=1)

            cols = [
                "Tipo", "Favorecido", "CPF/CNPJ", "PF/PJ", "Banco", "Valor", "Data", "Hora", "Confiança",
                "Status", "Origem do nome", "Arquivo",
            ]
            self.tree_comp = ttk.Treeview(frame_tbl, columns=cols, show="headings")
            for c in cols:
                self.tree_comp.heading(c, text=c, command=lambda cc=c: self._ordenar_tabela(self.tree_comp, cc))

            widths = {
                "Tipo": 80, "Favorecido": 220, "CPF/CNPJ": 120, "PF/PJ": 60, "Banco": 160,
                "Valor": 95, "Data": 90, "Hora": 70, "Confiança": 80, "Status": 90, "Origem do nome": 150, "Arquivo": 300,
            }
            for c in cols:
                self.tree_comp.column(c, width=widths[c], anchor="w")

            self.tree_comp.grid(row=0, column=0, sticky="nsew")
            ysb = ttk.Scrollbar(frame_tbl, orient="vertical", command=self.tree_comp.yview)
            xsb = ttk.Scrollbar(frame_tbl, orient="horizontal", command=self.tree_comp.xview)
            self.tree_comp.configure(yscroll=ysb.set, xscroll=xsb.set)
            ysb.grid(row=0, column=1, sticky="ns")
            xsb.grid(row=1, column=0, sticky="ew")

            self.tree_comp.tag_configure("odd", background="#F7F7F7")
            self.tree_comp.tag_configure("baixa", foreground="#B00020")
            self.tree_comp.tag_configure("media", foreground="#8A6D00")
            self.tree_comp.tag_configure("alta", foreground="#006400")
            self.tree_comp.tag_configure("sem_texto", background="#FFF4E5")
            self.tree_comp.tag_configure("falha", background="#FDECEC")

            self.menu_ctx = tk.Menu(self.root, tearoff=0)
            self.menu_ctx.add_command(label="Abrir PDF", command=self._ctx_abrir_pdf)
            self.menu_ctx.add_command(label="Abrir pasta do arquivo", command=self._ctx_abrir_pasta)
            self.menu_ctx.add_command(label="Copiar caminho", command=self._ctx_copiar_caminho)
            self.menu_ctx.add_command(label="Copiar dados da linha", command=self._ctx_copiar_linha)

            frame_det = ttk.LabelFrame(self.tab_comp, text="Detalhes", padding=8)
            frame_det.grid(row=0, column=1, sticky="nsew")
            frame_det.columnconfigure(1, weight=1)

            self.det_vars: dict[str, tk.StringVar] = {}
            campos = [
                "tipo", "favorecido", "documento", "tipo_pessoa", "banco", "valor", "data", "horario",
                "origem_nome", "confianca_num", "confianca_label", "status", "observacoes", "arquivo",
            ]
            for i, c in enumerate(campos):
                ttk.Label(frame_det, text=c.replace("_", " ").title() + ":").grid(row=i, column=0, sticky="nw", pady=1)
                v = tk.StringVar(value="-")
                self.det_vars[c] = v
                ttk.Label(frame_det, textvariable=v, wraplength=380).grid(row=i, column=1, sticky="nw", pady=1)

            frm_acoes = ttk.Frame(frame_det)
            frm_acoes.grid(row=len(campos), column=0, columnspan=2, sticky="ew", pady=(10, 0))
            ttk.Button(frm_acoes, text="Abrir PDF", command=self._ctx_abrir_pdf).pack(side="left", padx=2)
            ttk.Button(frm_acoes, text="Abrir pasta", command=self._ctx_abrir_pasta).pack(side="left", padx=2)
            ttk.Button(frm_acoes, text="Copiar caminho", command=self._ctx_copiar_caminho).pack(side="left", padx=2)
            ttk.Button(frm_acoes, text="Copiar resumo", command=self._copiar_resumo_item).pack(side="left", padx=2)

        def _build_tab_baixa(self) -> None:
            self.tab_baixa.columnconfigure(0, weight=1)
            self.tab_baixa.rowconfigure(1, weight=1)
            top = ttk.Frame(self.tab_baixa, padding=6)
            top.grid(row=0, column=0, sticky="ew")
            ttk.Button(top, text="Exportar baixa confiança", command=self._exportar_baixa).pack(side="left")
            self.tree_baixa = ttk.Treeview(self.tab_baixa, columns=("tipo", "favorecido", "banco", "valor", "conf", "arquivo"), show="headings")
            for c in ("tipo", "favorecido", "banco", "valor", "conf", "arquivo"):
                self.tree_baixa.heading(c, text=c.upper())
            self.tree_baixa.grid(row=1, column=0, sticky="nsew")

        def _build_tab_falhas(self) -> None:
            self.tab_falhas.columnconfigure(0, weight=1)
            self.tab_falhas.rowconfigure(1, weight=1)
            top = ttk.Frame(self.tab_falhas, padding=6)
            top.grid(row=0, column=0, sticky="ew")
            ttk.Button(top, text="Exportar falhas", command=self._exportar_falhas).pack(side="left")
            self.tree_falhas = ttk.Treeview(self.tab_falhas, columns=("arquivo", "erro", "tipo"), show="headings")
            for c, w in (("arquivo", 420), ("erro", 620), ("tipo", 150)):
                self.tree_falhas.heading(c, text=c.upper())
                self.tree_falhas.column(c, width=w, anchor="w")
            self.tree_falhas.grid(row=1, column=0, sticky="nsew")

        def _build_tab_revisao(self) -> None:
            self.tab_revisao.columnconfigure(0, weight=1)
            self.tab_revisao.rowconfigure(1, weight=1)
            topo = ttk.Frame(self.tab_revisao, padding=6)
            topo.grid(row=0, column=0, sticky="ew")
            ttk.Button(topo, text="Adicionar selecionado p/ revisão", command=self._marcar_revisao).pack(side="left")
            ttk.Button(topo, text="Marcar como revisado", command=self._marcar_revisado).pack(side="left", padx=4)
            self.txt_obs_revisao = ttk.Entry(topo)
            self.txt_obs_revisao.pack(side="left", fill="x", expand=True, padx=4)
            self.tree_revisao = ttk.Treeview(self.tab_revisao, columns=("arquivo", "favorecido", "status", "obs"), show="headings")
            for c in ("arquivo", "favorecido", "status", "obs"):
                self.tree_revisao.heading(c, text=c.upper())
            self.tree_revisao.grid(row=1, column=0, sticky="nsew")

        # ---------- Eventos e bindings ----------
        def _aplicar_bindings(self) -> None:
            self.root.bind("<Control-f>", lambda e: self.ent_busca.focus_set())
            self.root.bind("<F5>", lambda e: self._on_mapear())
            self.root.bind("<Escape>", lambda e: self._on_cancelar())
            self.tree_comp.bind("<<TreeviewSelect>>", self._on_select_item)
            self.tree_comp.bind("<Double-1>", lambda e: self._ctx_abrir_pdf())
            self.tree_comp.bind("<Button-3>", self._on_context_menu)
            for v in [self.var_busca, self.var_f_tipo, self.var_f_banco, self.var_f_conf, self.var_f_status, self.var_f_pessoa]:
                v.trace_add("write", lambda *_: self._aplicar_filtros())
            self.root.protocol("WM_DELETE_WINDOW", self._on_close)

        # ---------- Processamento em background ----------
        def _on_mapear(self) -> None:
            pasta = Path(self.var_pasta.get().strip())
            if not pasta.exists() or not pasta.is_dir():
                messagebox.showwarning("Pasta inválida", "Selecione uma pasta válida para iniciar a análise.")
                return
            if self.processando:
                return
            self.processando = True
            self.cancelar.clear()
            self.inicio_processamento = time.time()
            self.processados = 0
            self.total_alvo = len(list(pasta.rglob("*.pdf")))
            self.progress.configure(value=0, maximum=max(self.total_alvo, 1))
            self.var_progresso_texto.set("Iniciando análise...")
            self._log(f"Início do processamento de {self.total_alvo} PDF(s) em {pasta}")
            self._aplicar_estado_botoes()

            self.worker = threading.Thread(target=self._worker_mapear, args=(pasta,), daemon=True)
            self.worker.start()

        def _worker_mapear(self, pasta: Path) -> None:
            try:
                if self.cancelar.is_set():
                    self.fila_ui.put(("cancelado", None))
                    return
                comps, resumo, falhas = mapear_pasta_e_gerar_resumo(pasta, recursivo=True)
                self.fila_ui.put(("resultado", (comps, resumo, falhas)))
            except Exception as exc:
                self.fila_ui.put(("erro", str(exc)))

        def _poll_fila(self) -> None:
            try:
                while True:
                    tipo, payload = self.fila_ui.get_nowait()
                    if tipo == "resultado":
                        comps, resumo, falhas = payload
                        self.dataset_bruto = comps
                        self.falhas = falhas
                        self.ultimo_resumo = resumo
                        self._reconstruir_filtros_dinamicos()
                        self._aplicar_filtros()
                        self._atualizar_cards()
                        self._atualizar_resumo_tab()
                        self._atualizar_falhas_tab()
                        self._atualizar_baixa_tab()
                        self._log(f"Fim do processamento: {len(comps)} item(ns), {len(falhas)} falha(s).")
                        self._oferecer_pos_processamento()
                        self.var_status.set("Inventário concluído com sucesso")
                        self.var_progresso_texto.set(f"Concluído em {int(time.time() - self.inicio_processamento)}s")
                        self.progress.configure(value=self.total_alvo)
                        self.processando = False
                        self._aplicar_estado_botoes()
                    elif tipo == "erro":
                        self.processando = False
                        self._aplicar_estado_botoes()
                        self.var_status.set("Erro durante o processamento")
                        self._erro_amigavel("Falha durante a análise da pasta.", RuntimeError(str(payload)))
                        self._log(f"Erro: {payload}")
                    elif tipo == "cancelado":
                        self.processando = False
                        self._aplicar_estado_botoes()
                        self.var_status.set("Processamento cancelado pelo usuário")
                        self._log("Processamento cancelado pelo usuário.")
            except queue.Empty:
                pass
            self.root.after(200, self._poll_fila)

        def _on_cancelar(self) -> None:
            if self.processando:
                self.cancelar.set()
                self.processando = False
                self._aplicar_estado_botoes()
                self.var_status.set("Cancelando...")
                self._log("Solicitação de cancelamento enviada.")

        # ---------- Filtros / ordenação ----------
        def _aplicar_filtros(self) -> None:
            q = self.var_busca.get().strip().lower()
            f_tipo = self.var_f_tipo.get()
            f_banco = self.var_f_banco.get()
            f_conf = self.var_f_conf.get()
            f_status = self.var_f_status.get()
            f_pessoa = self.var_f_pessoa.get()

            def ok(c: Comprovante) -> bool:
                if f_tipo != "Todos" and c.tipo_comprovante != f_tipo:
                    return False
                if f_banco != "Todos" and c.banco != f_banco:
                    return False
                if f_conf != "Todos" and c.confianca_label != f_conf:
                    return False
                if f_status != "Todos" and c.status_extracao != f_status:
                    return False
                if f_pessoa != "Todos" and c.tipo_pessoa_favorecida != f_pessoa:
                    return False
                if q:
                    bag = " ".join([c.recebedor, c.documento_favorecido, c.banco, str(c.arquivo_pdf), c.tipo_comprovante]).lower()
                    if q not in bag:
                        return False
                return True

            self.dataset_filtrado = [c for c in self.dataset_bruto if ok(c)]
            self.var_total_filtrado.set(f"{len(self.dataset_filtrado)} / {len(self.dataset_bruto)}")
            self._render_tabela_principal()

        def _ordenar_tabela(self, tree: ttk.Treeview, coluna: str) -> None:
            if tree is not self.tree_comp:
                return
            key_map = {
                "Tipo": lambda c: c.tipo_comprovante,
                "Favorecido": lambda c: c.recebedor,
                "CPF/CNPJ": lambda c: c.documento_favorecido,
                "PF/PJ": lambda c: c.tipo_pessoa_favorecida,
                "Banco": lambda c: c.banco,
                "Valor": lambda c: valor_comprovante(c),
                "Data": lambda c: c.data_pagamento,
                "Hora": lambda c: c.horario_pagamento,
                "Confiança": lambda c: c.confianca_extracao,
                "Status": lambda c: c.status_extracao,
                "Origem do nome": lambda c: c.origem_nome_extraido,
                "Arquivo": lambda c: str(c.arquivo_pdf),
            }
            if coluna not in key_map:
                return
            if self.coluna_ordenacao == coluna:
                self.ordem_desc = not self.ordem_desc
            else:
                self.coluna_ordenacao = coluna
                self.ordem_desc = False
            self.dataset_filtrado.sort(key=key_map[coluna], reverse=self.ordem_desc)
            self._render_tabela_principal()

        # ---------- Atualizações de UI ----------
        def _render_tabela_principal(self) -> None:
            self.tree_comp.delete(*self.tree_comp.get_children())
            for i, c in enumerate(self.dataset_filtrado):
                tags = ["odd" if i % 2 else ""]
                tags.append(c.confianca_label)
                tags.append(c.status_extracao)
                self.tree_comp.insert(
                    "", "end", iid=str(c.arquivo_pdf),
                    values=(
                        c.tipo_comprovante, c.recebedor, c.documento_favorecido, c.tipo_pessoa_favorecida,
                        c.banco, formatar_moeda_br(valor_comprovante(c)), c.data_pagamento, c.horario_pagamento,
                        c.confianca_extracao, c.status_extracao, c.origem_nome_extraido, str(c.arquivo_pdf),
                    ),
                    tags=[t for t in tags if t],
                )

        def _atualizar_cards(self) -> None:
            r = self.ultimo_resumo
            if not r:
                return
            self.cards["pdfs"].set(str(r.total_pdfs))
            self.cards["proc"].set(str(r.total_processados))
            self.cards["valor"].set(formatar_moeda_br(r.soma_total_valores))
            self.cards["baixa"].set(str(r.total_baixa_confianca))
            self.cards["sem"].set(str(r.total_sem_texto))
            self.cards["falhas"].set(str(r.total_falhas))
            perc_baixa = (r.total_baixa_confianca / r.total_processados * 100) if r.total_processados else 0
            self.cards["%baixa"].set(f"{perc_baixa:.1f}%")
            dominante = "-"
            if r.por_banco:
                dominante = max(r.por_banco.items(), key=lambda x: x[1])[0]
            self.cards["dominante"].set(dominante)
            self.cards["valor_boletos"].set(formatar_moeda_br(r.soma_valores_boletos))
            self.cards["boletos_validos"].set(str(r.quantidade_boletos_com_valor_valido))

        def _atualizar_resumo_tab(self) -> None:
            r = self.ultimo_resumo
            if not r:
                return
            valores_validos = [valor_comprovante(c) for c in self.dataset_bruto if valor_comprovante(c) > 0]
            maior = max(valores_validos, default=0)
            menor = min(valores_validos, default=0)
            media = (r.soma_total_valores / len(valores_validos)) if valores_validos else 0

            tipo_dom = "-"
            tipo_count = Counter(c.tipo_comprovante for c in self.dataset_bruto)
            if tipo_count:
                tipo_dom = tipo_count.most_common(1)[0][0]

            def pct(v, t):
                return f"{(v/t*100):.1f}%" if t else "0.0%"

            linhas = [
                f"Pasta analisada: {r.pasta_analisada}",
                f"Total PDFs: {r.total_pdfs}",
                f"Total processado: {r.total_processados}",
                f"Sem texto: {r.total_sem_texto} ({pct(r.total_sem_texto, r.total_pdfs)})",
                f"Falhas: {r.total_falhas} ({pct(r.total_falhas, r.total_pdfs)})",
                f"Baixa confiança: {r.total_baixa_confianca} ({pct(r.total_baixa_confianca, r.total_processados)})",
                f"Valor total: {formatar_moeda_br(r.soma_total_valores)}",
                f"Valor total dos boletos: {formatar_moeda_br(r.soma_valores_boletos)}",
                f"Valor total do pix: {formatar_moeda_br(r.soma_valores_pix)}",
                f"Valor total das transferências: {formatar_moeda_br(r.soma_valores_transferencias)}",
                f"Valor total dos desconhecidos: {formatar_moeda_br(r.soma_valores_desconhecidos)}",
                f"Valor médio: {formatar_moeda_br(media)}",
                f"Maior valor: {formatar_moeda_br(maior)}",
                f"Menor valor: {formatar_moeda_br(menor)}",
                f"Por tipo -> Boleto: {r.quantidade_boletos} | Pix: {r.quantidade_pix} | Transferência: {r.quantidade_transferencias} | Desconhecido: {r.quantidade_desconhecidos}",
                f"PF/PJ -> PF: {r.total_pf} | PJ: {r.total_pj}",
                f"Banco mais recorrente: {self.cards['dominante'].get()}",
                f"Tipo mais recorrente: {tipo_dom}",
            ]
            self.txt_resumo.configure(state="normal")
            self.txt_resumo.delete("1.0", "end")
            self.txt_resumo.insert("end", "\n".join(linhas))
            self.txt_resumo.configure(state="disabled")

        def _atualizar_baixa_tab(self) -> None:
            self.dataset_baixa = [c for c in self.dataset_bruto if c.confianca_label == "baixa"]
            self.tree_baixa.delete(*self.tree_baixa.get_children())
            for c in self.dataset_baixa:
                self.tree_baixa.insert("", "end", values=(c.tipo_comprovante, c.recebedor, c.banco, c.valor, c.confianca_extracao, str(c.arquivo_pdf)))

        def _atualizar_falhas_tab(self) -> None:
            self.tree_falhas.delete(*self.tree_falhas.get_children())
            for f in self.falhas:
                erro = f.erro
                tipo = "leitura"
                if "permission" in erro.lower():
                    tipo = "permissao"
                elif "pdf" in erro.lower():
                    tipo = "pdf"
                self.tree_falhas.insert("", "end", values=(str(f.arquivo), erro, tipo))

        def _reconstruir_filtros_dinamicos(self) -> None:
            def setvals(cb, vals):
                atual = cb.get() or "Todos"
                cb["values"] = ["Todos"] + sorted(v for v in vals if v)
                cb.set(atual if atual in cb["values"] else "Todos")

            setvals(self.cb_tipo, {c.tipo_comprovante for c in self.dataset_bruto})
            setvals(self.cb_banco, {c.banco for c in self.dataset_bruto})
            setvals(self.cb_conf, {c.confianca_label for c in self.dataset_bruto})
            setvals(self.cb_status, {c.status_extracao for c in self.dataset_bruto})
            setvals(self.cb_pessoa, {c.tipo_pessoa_favorecida for c in self.dataset_bruto})

        def _atualizar_detalhes(self, c: Comprovante | None) -> None:
            if not c:
                for v in self.det_vars.values():
                    v.set("-")
                return
            self.det_vars["tipo"].set(c.tipo_comprovante)
            self.det_vars["favorecido"].set(c.recebedor)
            self.det_vars["documento"].set(c.documento_favorecido)
            self.det_vars["tipo_pessoa"].set(c.tipo_pessoa_favorecida)
            self.det_vars["banco"].set(c.banco)
            self.det_vars["valor"].set(formatar_moeda_br(valor_comprovante(c)))
            self.det_vars["data"].set(c.data_pagamento)
            self.det_vars["horario"].set(c.horario_pagamento)
            self.det_vars["origem_nome"].set(c.origem_nome_extraido)
            self.det_vars["confianca_num"].set(c.confianca_extracao)
            self.det_vars["confianca_label"].set(c.confianca_label)
            self.det_vars["status"].set(c.status_extracao)
            self.det_vars["observacoes"].set(c.observacoes_extracao)
            self.det_vars["arquivo"].set(str(c.arquivo_pdf))

        def _on_select_item(self, _event=None) -> None:
            sel = self.tree_comp.selection()
            if not sel:
                self.item_selecionado = None
                self._atualizar_detalhes(None)
                return
            iid = sel[0]
            for c in self.dataset_filtrado:
                if str(c.arquivo_pdf) == iid:
                    self.item_selecionado = c
                    self._atualizar_detalhes(c)
                    break

        # ---------- Ações ----------
        def _on_selecionar_pasta(self) -> None:
            p = selecionar_pasta_tk()
            if p:
                self.var_pasta.set(str(p))
                self._validar_pasta_ui()

        def _on_abrir_pasta(self) -> None:
            p = Path(self.var_pasta.get().strip())
            if p.exists():
                abrir_pasta(p)
            else:
                messagebox.showwarning("Pasta inválida", "Selecione uma pasta válida primeiro.")

        def _on_abrir_excel(self) -> None:
            if self.ultimo_excel and self.ultimo_excel.exists():
                try:
                    if os.name == "nt":
                        os.startfile(str(self.ultimo_excel))  # type: ignore[attr-defined]
                    elif sys.platform == "darwin":
                        subprocess.run(["open", str(self.ultimo_excel)], check=False)
                    else:
                        subprocess.run(["xdg-open", str(self.ultimo_excel)], check=False)
                except Exception as exc:
                    self._erro_amigavel("Não foi possível abrir o arquivo Excel.", exc)
            else:
                messagebox.showinfo("Excel", "Nenhum Excel disponível para abrir.")

        def _on_exportar_excel(self) -> None:
            p = Path(self.var_pasta.get().strip())
            if not p.exists() or not p.is_dir():
                messagebox.showwarning("Pasta inválida", "Selecione uma pasta válida antes de exportar.")
                return
            try:
                if self.dataset_bruto:
                    # Exporta dataset atual
                    resumo = self.ultimo_resumo or gerar_resumo(p, self.dataset_bruto, self.falhas, len(self.dataset_bruto), 0)
                    destino = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="relatorio_comprovantes.xlsx")
                    if not destino:
                        return
                    self.ultimo_excel = exportar_excel_completo(self.dataset_bruto, resumo, self.falhas, Path(destino))
                else:
                    self.ultimo_excel = mapear_pasta_e_alimentar_excel(p, recursivo=True)
                self._log(f"Exportação concluída: {self.ultimo_excel}")
                self.var_status.set(f"Excel gerado: {self.ultimo_excel}")
                messagebox.showinfo("Exportação", f"Excel gerado em:\n{self.ultimo_excel}")
                self._aplicar_estado_botoes()
            except Exception as exc:
                self._erro_amigavel("Não foi possível gerar o Excel.", exc)

        def _processar_lote_arquivo(self, lote: Path) -> None:
            if self.processando:
                messagebox.showinfo("Processamento", "Aguarde o processamento atual terminar.")
                return
            if not lote.exists() or lote.suffix.lower() != ".pdf":
                messagebox.showwarning("PDF inválido", f"Arquivo inválido:\n{lote}")
                return

            try:
                self.processando = True
                self._aplicar_estado_botoes()
                self.var_status.set("Processando PDF lote...")
                self.var_progresso_texto.set(f"Lendo {lote.name}")
                self._log(f"Início extração de lote: {lote}")

                texto_inicial = ler_texto_primeira_pagina(lote)
                if not texto_inicial:
                    raise ValueError("Não foi possível ler texto da primeira página. Use OCR no PDF.")

                base = extrair_dados_comprovante(texto_inicial)
                destino = criar_pasta_destino(base["banco"], base["data_pagamento"])
                comps, sem_texto = dividir_pdf_e_salvar(lote, destino)
                resumo = gerar_resumo(destino, comps, [], len(comps), sem_texto)
                excel = destino / "relatorio_comprovantes.xlsx"
                exportar_excel_completo(comps, resumo, [], excel)

                self.ultimo_excel = excel
                self.var_pasta.set(str(destino))
                self._validar_pasta_ui()

                # Atualiza dashboard com os novos dados extraídos do lote
                self.dataset_bruto = comps
                self.falhas = []
                self.ultimo_resumo = resumo
                self._reconstruir_filtros_dinamicos()
                self._aplicar_filtros()
                self._atualizar_cards()
                self._atualizar_resumo_tab()
                self._atualizar_baixa_tab()
                self._atualizar_falhas_tab()

                self.progress.configure(value=max(len(comps), 1), maximum=max(len(comps), 1))
                self.var_progresso_texto.set(f"Lote concluído: {len(comps)} comprovante(s)")
                self.var_status.set(f"Lote processado com sucesso | Excel: {excel.name}")
                self._log(f"Lote concluído. Pasta: {destino} | Excel: {excel}")

                messagebox.showinfo(
                    "Processamento concluído",
                    f"Comprovantes processados: {len(comps)}\n"
                    f"Arquivos sem texto: {sem_texto}\n"
                    f"Pasta de saída: {destino}\n"
                    f"Excel gerado: {excel}\n\n"
                    "Próximos passos: abrir Excel, abrir pasta ou iniciar nova análise.",
                )
            except Exception as exc:
                self.var_status.set("Erro no processamento do lote")
                self._erro_amigavel("Falha ao processar PDF lote.", exc)
            finally:
                self.processando = False
                self._aplicar_estado_botoes()

        def _on_processar_lote_manual(self) -> None:
            lote = selecionar_arquivo_pdf_manual()
            if not lote:
                self._log("Extração de lote cancelada pelo usuário.")
                return
            self._processar_lote_arquivo(lote)

        def _exportar_baixa(self) -> None:
            if not self.dataset_baixa:
                messagebox.showinfo("Baixa confiança", "Não há registros de baixa confiança para exportar.")
                return
            p = Path(self.var_pasta.get().strip()) if self.var_pasta.get().strip() else caminho_base_usuario()
            out = p / f"relatorio_baixa_confianca_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            resumo = self.ultimo_resumo or gerar_resumo(p, self.dataset_baixa, [], len(self.dataset_baixa), 0)
            exportar_excel_completo(self.dataset_baixa, resumo, [], out)
            self._log(f"Exportado baixa confiança: {out}")
            messagebox.showinfo("Exportação", f"Arquivo gerado:\n{out}")

        def _exportar_falhas(self) -> None:
            if not self.falhas:
                messagebox.showinfo("Falhas", "Não há falhas para exportar.")
                return
            p = Path(self.var_pasta.get().strip()) if self.var_pasta.get().strip() else caminho_base_usuario()
            out = p / f"falhas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
            lines = ["arquivo;erro\n"] + [f"{f.arquivo};{f.erro}\n" for f in self.falhas]
            out.write_text("".join(lines), encoding="utf-8")
            self._log(f"Falhas exportadas: {out}")
            messagebox.showinfo("Exportação", f"Falhas exportadas em:\n{out}")

        def _on_limpar_filtros(self) -> None:
            self.var_busca.set("")
            self.var_f_tipo.set("Todos")
            self.var_f_banco.set("Todos")
            self.var_f_conf.set("Todos")
            self.var_f_status.set("Todos")
            self.var_f_pessoa.set("Todos")
            self._aplicar_filtros()

        # ---------- Context menu ----------
        def _on_context_menu(self, event) -> None:
            iid = self.tree_comp.identify_row(event.y)
            if iid:
                self.tree_comp.selection_set(iid)
                self._on_select_item()
                self.menu_ctx.tk_popup(event.x_root, event.y_root)

        def _ctx_abrir_pdf(self) -> None:
            if self.item_selecionado and self.item_selecionado.arquivo_pdf.exists():
                try:
                    alvo = self.item_selecionado.arquivo_pdf
                    if os.name == "nt":
                        os.startfile(str(alvo))  # type: ignore[attr-defined]
                    elif sys.platform == "darwin":
                        subprocess.run(["open", str(alvo)], check=False)
                    else:
                        subprocess.run(["xdg-open", str(alvo)], check=False)
                except Exception as exc:
                    self._erro_amigavel("Não foi possível abrir o PDF selecionado.", exc)
            else:
                messagebox.showwarning("Arquivo", "Nenhum arquivo selecionado ou arquivo inexistente.")

        def _ctx_abrir_pasta(self) -> None:
            if self.item_selecionado:
                abrir_pasta(self.item_selecionado.arquivo_pdf.parent)

        def _ctx_copiar_caminho(self) -> None:
            if not self.item_selecionado:
                return
            self.root.clipboard_clear()
            self.root.clipboard_append(str(self.item_selecionado.arquivo_pdf))
            self._log("Caminho copiado para a área de transferência.")

        def _ctx_copiar_linha(self) -> None:
            if not self.item_selecionado:
                return
            c = self.item_selecionado
            txt = " | ".join([
                c.tipo_comprovante, c.recebedor, c.documento_favorecido, c.tipo_pessoa_favorecida,
                c.banco, c.valor, c.data_pagamento, c.horario_pagamento, c.confianca_extracao,
                c.status_extracao, str(c.arquivo_pdf),
            ])
            self.root.clipboard_clear()
            self.root.clipboard_append(txt)
            self._log("Dados da linha copiados.")

        def _copiar_resumo_item(self) -> None:
            if not self.item_selecionado:
                return
            c = self.item_selecionado
            txt = (
                f"Tipo: {c.tipo_comprovante}\n"
                f"Favorecido: {c.recebedor}\n"
                f"Documento: {c.documento_favorecido}\n"
                f"Banco: {c.banco}\n"
                f"Valor: {formatar_moeda_br(valor_comprovante(c))}\n"
                f"Data/Hora: {c.data_pagamento} {c.horario_pagamento}\n"
                f"Confiança: {c.confianca_extracao} ({c.confianca_label})\n"
                f"Arquivo: {c.arquivo_pdf}"
            )
            self.root.clipboard_clear()
            self.root.clipboard_append(txt)
            self._log("Resumo do registro copiado.")

        # ---------- Revisão manual ----------
        def _marcar_revisao(self) -> None:
            if not self.item_selecionado:
                messagebox.showinfo("Revisão", "Selecione um item na aba Comprovantes.")
                return
            key = str(self.item_selecionado.arquivo_pdf)
            self.revisao_manual[key] = {
                "revisado": False,
                "obs": self.txt_obs_revisao.get().strip(),
            }
            self._render_revisao()

        def _marcar_revisado(self) -> None:
            sel = self.tree_revisao.selection()
            if not sel:
                return
            key = sel[0]
            if key in self.revisao_manual:
                self.revisao_manual[key]["revisado"] = True
                if self.txt_obs_revisao.get().strip():
                    self.revisao_manual[key]["obs"] = self.txt_obs_revisao.get().strip()
            self._render_revisao()

        def _render_revisao(self) -> None:
            self.tree_revisao.delete(*self.tree_revisao.get_children())
            by_path = {str(c.arquivo_pdf): c for c in self.dataset_bruto}
            for path, meta in self.revisao_manual.items():
                c = by_path.get(path)
                fav = c.recebedor if c else "-"
                status = "revisado" if meta.get("revisado") else "pendente"
                obs = str(meta.get("obs", ""))
                self.tree_revisao.insert("", "end", iid=path, values=(path, fav, status, obs))

        # ---------- Utilitários UI ----------
        def _validar_pasta_ui(self) -> None:
            ptxt = self.var_pasta.get().strip()
            p = Path(ptxt) if ptxt else None
            if p and p.exists() and p.is_dir():
                self.var_pasta_ok.set("Pasta válida")
            else:
                self.var_pasta_ok.set("Pasta inválida")
            self._aplicar_estado_botoes()

        def _aplicar_estado_botoes(self) -> None:
            ptxt = self.var_pasta.get().strip()
            pasta_ok = bool(ptxt and Path(ptxt).exists() and Path(ptxt).is_dir())
            self.btn_mapear.configure(state="disabled" if (not pasta_ok or self.processando) else "normal")
            self.btn_atualizar.configure(state="disabled" if (not pasta_ok or self.processando) else "normal")
            self.btn_exportar.configure(state="disabled" if (not pasta_ok or self.processando) else "normal")
            self.btn_abrir_pasta.configure(state="disabled" if not pasta_ok else "normal")
            self.btn_cancelar.configure(state="normal" if self.processando else "disabled")
            self.btn_abrir_excel.configure(state="normal" if (self.ultimo_excel and Path(self.ultimo_excel).exists()) else "disabled")

        def _oferecer_pos_processamento(self) -> None:
            opcoes = (
                "Processamento concluído com sucesso.\n\n"
                "Deseja abrir o arquivo Excel agora?\n"
                "(Se escolher Não, você ainda poderá abrir a pasta de saída.)"
            )
            try:
                from tkinter import messagebox
                if self.ultimo_excel and self.ultimo_excel.exists() and messagebox.askyesno("Concluído", opcoes):
                    self._on_abrir_excel()
                    return
                if messagebox.askyesno("Concluído", "Deseja abrir a pasta de saída?"):
                    self._on_abrir_pasta()
            except Exception:
                pass

        def _erro_amigavel(self, titulo: str, exc: Exception, detalhe: str = "") -> None:
            msg = f"{titulo}\n\n{exc}" if not detalhe else f"{titulo}\n\n{detalhe}\n\n{exc}"
            self._log(f"ERRO: {msg}")
            try:
                import tkinter as tk
                from tkinter import messagebox
                if messagebox.askyesno("Erro", msg + "\n\nDeseja copiar detalhes para suporte?"):
                    self.root.clipboard_clear()
                    self.root.clipboard_append(msg)
                if messagebox.askyesno("Logs", "Deseja abrir a pasta de logs agora?"):
                    abrir_pasta(pasta_logs_app())
            except Exception:
                pass

        def _log(self, msg: str) -> None:
            stamp = datetime.now().strftime("%H:%M:%S")
            self.txt_log.configure(state="normal")
            self.txt_log.insert("end", f"[{stamp}] {msg}\n")
            self.txt_log.see("end")
            self.txt_log.configure(state="disabled")
            try:
                self.logger.info(msg)
            except Exception:
                pass

        # ---------- Persistência ----------
        def _carregar_config(self) -> None:
            try:
                if not self.CONFIG_PATH.exists():
                    return
                cfg = json.loads(self.CONFIG_PATH.read_text(encoding="utf-8"))
                self.var_pasta.set(cfg.get("ultima_pasta", ""))
                self.var_f_tipo.set(cfg.get("f_tipo", "Todos"))
                self.var_f_banco.set(cfg.get("f_banco", "Todos"))
                self.var_f_conf.set(cfg.get("f_conf", "Todos"))
                self.var_f_status.set(cfg.get("f_status", "Todos"))
                self.var_f_pessoa.set(cfg.get("f_pessoa", "Todos"))
                self.var_busca.set(cfg.get("busca", ""))
                if "geometry" in cfg:
                    self.root.geometry(cfg["geometry"])
                if "aba" in cfg:
                    try:
                        self.notebook.select(cfg["aba"])
                    except Exception:
                        pass
                self.ultimo_excel = Path(cfg["ultimo_excel"]) if cfg.get("ultimo_excel") else None
                self._validar_pasta_ui()
            except Exception:
                pass

        def _salvar_config(self) -> None:
            cfg = {
                "ultima_pasta": self.var_pasta.get().strip(),
                "f_tipo": self.var_f_tipo.get(),
                "f_banco": self.var_f_banco.get(),
                "f_conf": self.var_f_conf.get(),
                "f_status": self.var_f_status.get(),
                "f_pessoa": self.var_f_pessoa.get(),
                "busca": self.var_busca.get(),
                "geometry": self.root.geometry(),
                "aba": self.notebook.index("current"),
                "ultimo_excel": str(self.ultimo_excel) if self.ultimo_excel else "",
            }
            # salva largura das colunas da tabela principal
            cfg["colunas"] = {c: self.tree_comp.column(c, "width") for c in self.tree_comp["columns"]}
            try:
                self.CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
            except Exception:
                pass

        def _on_close(self) -> None:
            self._salvar_config()
            self.root.destroy()

    root = tk.Tk()
    app = AppDashboard(
        root,
        pasta_inicial=pasta_inicial,
        lote_inicial=lote_inicial,
        auto_mapear=auto_mapear,
        auto_lote=auto_lote,
    )
    app._validar_pasta_ui()
    root.mainloop()


def executar_lancador(acao_lote, acao_inventario, acao_dashboard, acao_diagnostico) -> bool:
    """Abre um lançador GUI mais completo para uso principal no executável."""
    try:
        import tkinter as tk
        from tkinter import messagebox, ttk
    except Exception:
        return False

    root = tk.Tk()
    root.title(f"{APP_NOME} - Launcher")
    root.geometry("940x640")
    root.minsize(860, 560)

    style = ttk.Style(root)
    try:
        style.theme_use("vista")
    except Exception:
        pass
    style.configure("Title.TLabel", font=("Segoe UI", 16, "bold"))
    style.configure("Subtitle.TLabel", font=("Segoe UI", 10))
    style.configure("Primary.TButton", font=("Segoe UI", 10, "bold"), padding=(10, 8))

    frame = ttk.Frame(root, padding=16)
    frame.pack(fill="both", expand=True)
    frame.columnconfigure(0, weight=1)
    frame.rowconfigure(3, weight=1)

    ttk.Label(frame, text=APP_NOME, style="Title.TLabel").grid(row=0, column=0, sticky="w")
    ttk.Label(
        frame,
        text=(
            "Aplicativo desktop para organizar comprovantes em PDF, processar lotes, "
            "inventariar pastas e gerar relatórios Excel com qualidade operacional."
        ),
        style="Subtitle.TLabel",
        wraplength=880,
        justify="left",
    ).grid(row=1, column=0, sticky="w", pady=(4, 10))

    painel = ttk.Frame(frame)
    painel.grid(row=2, column=0, sticky="nsew", pady=(0, 8))
    painel.columnconfigure(0, weight=2)
    painel.columnconfigure(1, weight=3)

    # Ações principais
    card_acoes = ttk.LabelFrame(painel, text="Ações principais", padding=12)
    card_acoes.grid(row=0, column=0, sticky="nsew", padx=(0, 8))

    def _rodar(callback):
        root.destroy()
        callback()

    ttk.Button(card_acoes, text="Abrir Dashboard", style="Primary.TButton", command=lambda: _rodar(acao_dashboard)).pack(fill="x", pady=4)
    ttk.Button(card_acoes, text="Processar PDF Lote", command=lambda: _rodar(acao_lote)).pack(fill="x", pady=4)
    ttk.Button(card_acoes, text="Inventariar Pasta", command=lambda: _rodar(acao_inventario)).pack(fill="x", pady=4)
    ttk.Button(card_acoes, text="Diagnóstico do Ambiente", command=lambda: _rodar(acao_diagnostico)).pack(fill="x", pady=4)

    # Diagnóstico visual
    card_diag = ttk.LabelFrame(painel, text="Status do ambiente", padding=12)
    card_diag.grid(row=0, column=1, sticky="nsew")
    card_diag.columnconfigure(0, weight=1)
    txt_diag = tk.Text(card_diag, height=14, wrap="word")
    txt_diag.grid(row=0, column=0, sticky="nsew")

    barra = ttk.Frame(card_diag)
    barra.grid(row=1, column=0, sticky="ew", pady=(8, 0))
    for i in range(5):
        barra.columnconfigure(i, weight=1)

    def _recarregar_diag() -> dict[str, object]:
        amb = verificar_ambiente_execucao()
        txt_diag.configure(state="normal")
        txt_diag.delete("1.0", "end")
        txt_diag.insert("end", formatar_diagnostico_texto(amb))
        txt_diag.configure(state="disabled")
        return amb

    def _copiar_diag() -> None:
        amb = verificar_ambiente_execucao()
        root.clipboard_clear()
        root.clipboard_append(formatar_diagnostico_texto(amb))
        messagebox.showinfo("Launcher", "Diagnóstico copiado para a área de transferência.")

    ttk.Button(barra, text="Recarregar", command=_recarregar_diag).grid(row=0, column=0, padx=3, sticky="ew")
    ttk.Button(barra, text="Copiar diagnóstico", command=_copiar_diag).grid(row=0, column=1, padx=3, sticky="ew")
    ttk.Button(barra, text="Abrir logs", command=lambda: abrir_pasta(pasta_logs_app())).grid(row=0, column=2, padx=3, sticky="ew")
    ttk.Button(barra, text="Abrir temp", command=lambda: abrir_pasta(pasta_temp_app())).grid(row=0, column=3, padx=3, sticky="ew")
    ttk.Button(barra, text="Base do usuário", command=lambda: abrir_pasta(caminho_base_usuario())).grid(row=0, column=4, padx=3, sticky="ew")

    _recarregar_diag()

    rodape = ttk.Frame(frame)
    rodape.grid(row=3, column=0, sticky="ew")
    ttk.Label(rodape, text=f"Versão {APP_VERSAO}").pack(side="left")
    ttk.Button(rodape, text="Fechar", command=root.destroy).pack(side="right")

    root.mainloop()
    return True


def diagnostico_ambiente() -> None:
    """Mostra contexto de execução para comparar VS Code vs .exe empacotado."""
    print("=== DIAGNÓSTICO DE AMBIENTE ===")
    print(f"Python executável: {sys.executable}")
    print(f"Versão Python: {sys.version.split()[0]}")
    print(f"Congelado (PyInstaller): {em_modo_frozen()}")
    print(f"Diretório atual: {Path.cwd()}")
    print(f"Base de execução: {diretorio_execucao_base()}")
    print(f"Base do usuário: {caminho_base_usuario()}")
    print(f"Arquivo principal: {Path(__file__).resolve()}")

    ambiente = verificar_ambiente_execucao()
    faltando = ambiente["faltando_modulos"]
    if faltando:
        print(f"[ERRO] Módulos ausentes: {', '.join(faltando)}")
    else:
        print("[OK] Módulos Python essenciais carregados.")

    print(f"[{'OK' if ambiente['tesseract'] else 'ERRO'}] tesseract: {ambiente['tesseract'] or 'não encontrado'}")
    print(f"[{'OK' if ambiente['pdftoppm'] else 'ERRO'}] pdftoppm: {ambiente['pdftoppm'] or 'não encontrado'}")

    if ambiente["avisos"]:
        print("\nAvisos de paridade VS Code x Executável:")
        for aviso in ambiente["avisos"]:
            print(f"- {aviso}")



def gerar_executavel_windows(script_path: Path, nome_exe: str, windowed: bool = False) -> int:
    """Gera executável via `py -m PyInstaller` usando o próprio script como entrada."""
    script_resolvido = script_path.resolve()
    if not script_resolvido.exists():
        print(f"[ERRO] Script não encontrado para build: {script_resolvido}")
        return 2

    dist_dir = script_resolvido.parent / "dist"
    work_dir = script_resolvido.parent / "build_pyinstaller"

    cmd = [
        sys.executable,
        "-m",
        "PyInstaller",
        "--onefile",
        "--clean",
        "--noconfirm",
        "--name",
        nome_exe,
        "--distpath",
        str(dist_dir),
        "--workpath",
        str(work_dir),
        "--specpath",
        str(script_resolvido.parent),
        "--runtime-tmpdir",
        str(pasta_temp_app()),
        "--hidden-import",
        "tkinter",
        "--hidden-import",
        "tkinter.ttk",
        "--hidden-import",
        "tkinter.filedialog",
        "--hidden-import",
        "tkinter.messagebox",
        "--hidden-import",
        "openpyxl",
        "--hidden-import",
        "pypdf",
        "--hidden-import",
        "PyPDF2",
        "--hidden-import",
        "pdf2image",
        "--hidden-import",
        "pytesseract",
        "--collect-all",
        "pdf2image",
        "--collect-all",
        "pytesseract",
    ]

    if windowed:
        cmd.append("--windowed")
    else:
        cmd.append("--console")

    cmd.append(str(script_resolvido))

    print("[BUILD] Comando:", " ".join(cmd))
    proc = subprocess.run(cmd, check=False, cwd=str(script_resolvido.parent))
    if proc.returncode == 0:
        print(f"[OK] Build concluído. Verifique: {script_resolvido.parent / 'dist' / (nome_exe + '.exe')}")
    else:
        print(f"[ERRO] Falha no build do executável (código {proc.returncode}).")
    return proc.returncode


# =========================
# CLI
# =========================

def argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Organiza comprovantes em PDF e gera Excel.")
    parser.add_argument("--lote", type=Path, help="PDF único com múltiplos comprovantes (um por página)")
    parser.add_argument("--inventario", type=Path, help="Pasta para inventariar PDFs")
    parser.add_argument("--excel", type=Path, default=None, help="Caminho do Excel final")
    parser.add_argument("--selecionar-lote", action="store_true")
    parser.add_argument("--selecionar-pasta", action="store_true")
    parser.add_argument("--dashboard", action="store_true", help="Abre o painel principal unificado")
    parser.add_argument("--launcher", action="store_true", help="Abre a tela inicial do aplicativo")
    parser.add_argument("--diagnostico-ambiente", action="store_true", help="Imprime diagnóstico para comparar VS Code e executável")
    parser.add_argument("--nao-abrir-pasta", action="store_true")
    parser.add_argument("--nao-recursivo", action="store_true")
    parser.add_argument("--gerar-exe", action="store_true", help="Gera executável .exe completo via PyInstaller")
    parser.add_argument("--gerar-exe-dashboard", action="store_true", help="Gera .exe dedicado ao dashboard unificado")
    parser.add_argument("--nome-exe", default="Organizador_Comprovantes", help="Nome do executável ao usar --gerar-exe")
    parser.add_argument("--exe-windowed", action="store_true", help="Gera .exe sem console (focado em dashboard)")
    return parser.parse_args()


def main() -> None:
    args = argumentos()

    if args.gerar_exe:
        rc = gerar_executavel_windows(Path(__file__), args.nome_exe, windowed=args.exe_windowed)
        if rc != 0:
            raise SystemExit(rc)
        return

    if args.gerar_exe_dashboard:
        nome_dashboard = f"{args.nome_exe}_Dashboard"
        rc = gerar_executavel_windows(Path(__file__), nome_dashboard, windowed=True)
        if rc != 0:
            raise SystemExit(rc)
        return

    if args.diagnostico_ambiente:
        diagnostico_ambiente()
        if not (args.dashboard or args.lote or args.inventario or args.selecionar_lote or args.selecionar_pasta):
            return

    def acao_lote() -> None:
        lote = args.lote or selecionar_arquivo_pdf_tk()
        executar_dashboard(
            pasta_inicial=None,
            lote_inicial=lote,
            auto_mapear=False,
            auto_lote=bool(lote),
        )

    def acao_inventario() -> None:
        pasta = args.inventario or selecionar_pasta_tk()
        executar_dashboard(
            pasta_inicial=pasta,
            lote_inicial=None,
            auto_mapear=bool(pasta),
            auto_lote=False,
        )

    def acao_dashboard() -> None:
        executar_dashboard(
            pasta_inicial=None,
            lote_inicial=None,
            auto_mapear=False,
            auto_lote=False,
        )

    sem_argumentos_operacionais = not any([
        args.dashboard,
        args.lote,
        args.inventario,
        args.selecionar_lote,
        args.selecionar_pasta,
    ])

    if args.launcher or (em_modo_frozen() and sem_argumentos_operacionais):
        if executar_lancador(acao_lote, acao_inventario, acao_dashboard, diagnostico_ambiente):
            return

    pasta_inicial = args.inventario or (selecionar_pasta_tk() if args.selecionar_pasta else None)
    lote_inicial = args.lote or (selecionar_arquivo_pdf_tk() if args.selecionar_lote else None)

    executar_dashboard(
        pasta_inicial=pasta_inicial,
        lote_inicial=lote_inicial,
        auto_mapear=bool(pasta_inicial),
        auto_lote=bool(lote_inicial),
    )


def tratar_erro_fatal(exc: Exception) -> None:
    logger = configurar_logging()
    logger.exception("Erro fatal: %s", exc)

    mensagem = (
        f"Ocorreu um erro inesperado.\n\n"
        f"{type(exc).__name__}: {exc}\n\n"
        f"Verifique o log em:\n"
        f"{pasta_logs_app()}"
    )

    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Erro inesperado", mensagem)
        root.destroy()
    except Exception:
        print(mensagem)


if __name__ == "__main__":
    try:
        import multiprocessing
        multiprocessing.freeze_support()
        main()
    except KeyboardInterrupt:
        raise SystemExit("Execução interrompida pelo usuário.")
    except Exception as exc:
        tratar_erro_fatal(exc)
        raise SystemExit(1)
